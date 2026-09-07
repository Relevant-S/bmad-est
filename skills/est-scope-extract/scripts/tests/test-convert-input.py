#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# ///
"""Tests for convert-input.py — normalization, citable anchors, graceful degradation.

The load-bearing property: every converted source carries anchors a citation can point at,
and a source that cannot be converted is reported rather than silently lost.
"""

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent.parent


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


convert = load("convert_input", "convert-input.py")


class TestConvert(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.out = self.dir / "normalized"
        self.out.mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    def test_csv_becomes_a_table_with_citable_row_numbers(self):
        src = self.dir / "backlog.csv"
        src.write_text("Feature,Priority\nLogin,Must\nExport,Should\n", encoding="utf-8")
        entry = convert.convert_one(src, self.out, 1)
        body = Path(entry["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("| row |", body)
        self.assertIn("| 2 | Login | Must |", body)

    def test_csv_semicolon_delimiter_is_sniffed(self):
        src = self.dir / "euro.csv"
        src.write_text("Feature;Priority\nLogin;Must\n", encoding="utf-8")
        body = Path(convert.convert_one(src, self.out, 1)["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("| 2 | Login | Must |", body)

    def test_pipe_in_source_data_is_escaped_not_dropped(self):
        src = self.dir / "piped.csv"
        src.write_text('Feature,Note\nLogin,"email|password"\n', encoding="utf-8")
        body = Path(convert.convert_one(src, self.out, 1)["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("email\\|password", body)

    def test_markdown_passes_through_unchanged(self):
        src = self.dir / "brief.md"
        src.write_text("# Brief\n\nUsers need accounts.\n", encoding="utf-8")
        entry = convert.convert_one(src, self.out, 3)
        self.assertEqual(entry["converter"], "passthrough")
        self.assertIn("Users need accounts.", Path(entry["converted_path"]).read_text(encoding="utf-8"))

    def test_source_header_records_provenance(self):
        src = self.dir / "brief.md"
        src.write_text("hello", encoding="utf-8")
        body = Path(convert.convert_one(src, self.out, 1)["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("<!-- source:", body)
        self.assertIn("converter: passthrough", body)

    def test_unsupported_extension_is_flagged_for_native_reading(self):
        src = self.dir / "diagram.sketch"
        src.write_bytes(b"\x00\x01")
        entry = convert.convert_one(src, self.out, 4)
        self.assertTrue(entry["needs_native_read"])
        self.assertIn("unsupported extension", entry["warning"])

    def test_missing_file_is_recorded_not_raised(self):
        entry = convert.convert_one(self.dir / "nope.docx", self.out, 5)
        self.assertTrue(entry["needs_native_read"])
        self.assertEqual(entry["warning"], "file not found")

    def test_html_is_stripped_to_text(self):
        src = self.dir / "page.html"
        src.write_text("<html><head><style>x{}</style></head><body><p>Scope: login</p></body></html>",
                       encoding="utf-8")
        body = Path(convert.convert_one(src, self.out, 6)["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("Scope: login", body)
        self.assertNotIn("x{}", body)

    @unittest.skipIf(convert._optional("openpyxl") is None, "openpyxl not installed")
    def test_multi_tab_workbook_keeps_sheet_names_as_anchors(self):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "Backlog"
        wb.active.append(["Feature", "Priority"])
        wb.active.append(["Login", "Must"])
        second = wb.create_sheet("Assumptions")
        second.append(["Assumption"])
        second.append(["Client provides the API"])
        src = self.dir / "book.xlsx"
        wb.save(src)

        body = Path(convert.convert_one(src, self.out, 7)["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("## sheet: Backlog", body)
        self.assertIn("## sheet: Assumptions", body)
        self.assertIn("Client provides the API", body)

    def test_csv_row_numbers_are_real_line_numbers(self):
        src = self.dir / "plain.csv"
        src.write_text("Feature,Priority\nLogin,Must\nExport,Should\n", encoding="utf-8")
        body = Path(convert.convert_one(src, self.out, 1)["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("| 1 | Feature | Priority |", body)
        self.assertIn("| 3 | Export | Should |", body)

    def test_quoted_multiline_field_survives_and_does_not_shift_later_rows(self):
        src = self.dir / "multiline.csv"
        src.write_text('Feature,Note\nLogin,"line one\nline two"\nExport,plain\n', encoding="utf-8")
        body = Path(convert.convert_one(src, self.out, 1)["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("line one<br>line two", body)   # not welded into "line oneline two"
        self.assertIn("| 4 | Export | plain |", body)  # real line 4, after the two-line field

    @unittest.skipIf(convert._optional("openpyxl") is None, "openpyxl not installed")
    def test_banner_and_blank_rows_do_not_shift_workbook_row_anchors(self):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "Backlog"
        wb.active.append(["ACME Backlog v3 — CONFIDENTIAL"])  # row 1
        wb.active.append([])                                   # row 2, blank
        wb.active.append(["Feature", "Priority"])              # row 3
        wb.active.append(["Login", "Must"])                    # row 4
        src = self.dir / "banner.xlsx"
        wb.save(src)

        body = Path(convert.convert_one(src, self.out, 1)["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("| 4 | Login | Must |", body)   # cited as row 4 because it IS row 4
        self.assertNotIn("| 3 | Login", body)
        self.assertNotIn("| 2 |", body)               # the blank row is skipped, not renumbered

    @unittest.skipIf(convert._optional("openpyxl") is None, "openpyxl not installed")
    def test_no_row_is_promoted_to_header(self):
        from openpyxl import Workbook

        wb = Workbook()
        wb.active.title = "Sheet1"
        wb.active.append(["Feature", "Priority"])
        wb.active.append(["Login", "Must"])
        src = self.dir / "noheader.xlsx"
        wb.save(src)

        body = Path(convert.convert_one(src, self.out, 1)["converted_path"]).read_text(encoding="utf-8")
        self.assertIn("| row | A | B |", body)          # generic labels; meaning left to the model
        self.assertIn("| 1 | Feature | Priority |", body)  # the real header stays a citable row

    def test_manifest_reports_unread_sources(self):
        (self.dir / "a.md").write_text("scope", encoding="utf-8")
        (self.dir / "b.sketch").write_bytes(b"\x00")
        sources = [convert.convert_one(p, self.out, i)
                   for i, p in enumerate(sorted(self.dir.glob("*.*")), start=1)]
        unread = [s["id"] for s in sources if s["needs_native_read"]]
        self.assertEqual(len(unread), 1)


if __name__ == "__main__":
    unittest.main()
