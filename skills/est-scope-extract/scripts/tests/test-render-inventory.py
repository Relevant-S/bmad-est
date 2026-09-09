#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Tests for render-inventory.py — the human and sales projections.

Two load-bearing properties. A verbatim quote must survive the round trip intact: pipes,
newlines and non-Latin text are exactly what a naive markdown table destroys, and destroying
the quote destroys the traceability the whole module rests on. And **every reference must
resolve** — the projections used to emit task ids matching no row in the file that named
them, dependencies as bare ids with their evidence dropped, and locations as prose. A
pointer nobody can follow is worse than none, because it reads as though it were checked.
"""

import contextlib
import csv
import importlib.util
import io
import json
import re
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from fixtures import feature, inventory  # noqa: E402

SCRIPTS = Path(__file__).resolve().parent.parent


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


render = load("render_inventory", "render-inventory.py")

try:
    import openpyxl  # noqa: F401
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def task(tid, name, location, quote):
    return {"id": tid, "name": name,
            "citations": [{"source_id": "S1", "location": location, "quote": quote}]}


ROW9 = ("Every data-access path applies the acting user's view state for the record's domain "
        "— Full, Context only or None — at the data layer, not by hiding UI.")
ROW10 = "Every data-access path applies the acting user's view state — Full or None."

NORMALIZED = f"""# Kampies-Post-discovery.xlsx

## sheet: Operators - Feature list

| row | epic | name | description |
| --- | --- | --- | --- |
| 9 | System | Enforce domain access | {ROW9} |
| 10 |  | Enforce domain access | {ROW10} |
"""


class TestRender(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, inv):
        path = self.dir / "feature-inventory.json"
        path.write_text(json.dumps(inv), encoding="utf-8")
        return path

    def test_markdown_preserves_a_quote_containing_pipes_and_newlines(self):
        hostile = "Users can pay by card | bank transfer\nand receive a receipt."
        f = feature()
        f["citations"][0]["quote"] = hostile
        md = render.render_markdown(inventory(features=[f]))
        self.assertIn("Users can pay by card | bank transfer", md)

    def test_markdown_preserves_original_language_quote(self):
        f = feature()
        f["citations"][0]["quote_original"] = "Користувачі повинні входити в систему."
        f["citations"][0]["quote_language"] = "uk"
        md = render.render_markdown(inventory(features=[f]))
        self.assertIn("Користувачі повинні входити в систему.", md)

    def test_markdown_carries_scope_and_says_nothing_about_cost(self):
        """These views render what the client asked for. The classification a story is priced
        on lives in classification.json and renders through est-estimate, so a reviewer reading
        this is checking the scope against the source with no effort judgement mixed in."""
        md = render.render_markdown(inventory())
        for axis in ("size_band", "compressibility", "review_tier", "clarity", "novelty"):
            self.assertNotIn(f"| {axis} |", md)
        self.assertIn("Users must be able to log in", md)

    def test_markdown_includes_not_scope_section(self):
        md = render.render_markdown(inventory())
        self.assertIn("Not treated as scope", md)
        self.assertIn("commercial terms, not scope", md)

    def test_markdown_marks_a_feature_outside_the_agreed_scope(self):
        md = render.render_markdown(inventory(features=[feature(scope_status="outside_agreed_scope")]))
        self.assertIn("outside agreed scope", md)
        self.assertIn("estimated separately", md)

    def test_csv_defaults_scope_status_when_absent(self):
        target = self.dir / "feature-inventory.csv"
        render.render_csv(inventory(), target)
        with target.open(encoding="utf-8-sig") as fh:
            row = next(iter(csv.DictReader(fh)))
        self.assertEqual(row["scope_status"], "in_agreed_scope")

    def test_csv_row_carries_the_scope_columns_and_no_cost_axis(self):
        target = self.dir / "feature-inventory.csv"
        render.render_csv(inventory(), target)
        with target.open(encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            row = next(iter(reader))
            columns = reader.fieldnames
        self.assertEqual(row["id"], "F1")
        self.assertEqual(row["commitment"], "committed")
        for axis in ("size_band", "compressibility", "review_tier", "clarity", "novelty"):
            self.assertNotIn(axis, columns)

    def test_csv_survives_a_quote_containing_a_comma_and_newline(self):
        f = feature()
        f["citations"][0]["quote"] = "Users can log in, log out\nand reset a password."
        target = self.dir / "feature-inventory.csv"
        render.render_csv(inventory(features=[f]), target)
        with target.open(encoding="utf-8-sig") as fh:
            row = next(iter(csv.DictReader(fh)))
        self.assertIn("reset a password", row["quotes"])

    def test_the_csv_carries_every_quote_not_only_the_first(self):
        """`primary_quote` published citations[0] and dropped the rest, so a story assembled
        from three rows shipped one of them."""
        f = feature()
        f["citations"].append({"source_id": "S1", "location": "§2.2",
                               "quote": "Sessions expire after 30 minutes."})
        target = self.dir / "feature-inventory.csv"
        render.render_csv(inventory(features=[f]), target)
        with target.open(encoding="utf-8-sig") as fh:
            row = next(iter(csv.DictReader(fh)))
        self.assertIn("able to log in", row["quotes"])
        self.assertIn("Sessions expire after 30 minutes", row["quotes"])

    def test_render_writes_both_projections(self):
        path = self.write(inventory())
        sys.argv = ["render-inventory.py", str(path), "--formats", "md,csv"]
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(render.main(), 0)
        self.assertTrue((self.dir / "feature-inventory.md").exists())
        self.assertTrue((self.dir / "feature-inventory.csv").exists())
        self.assertTrue((self.dir / "feature-inventory.tasks.csv").exists())


class Grouping(unittest.TestCase):
    """The counts a reader checks first: stories against epics against source rows."""

    def grouped(self):
        f = feature("F1", "Lead list: search, filter and sort", epic_id="E1",
                    surfaces=["backend", "frontend"],
                    tasks=[task("T1", "Search leads", "row 181", "Search the leads list by name."),
                           task("T2", "Filter leads", "row 182", "Filter leads by status.")])
        return inventory(features=[f],
                         epics=[{"id": "E1", "name": "Lead", "origin": "source"}])

    def test_the_ratio_that_shows_grouping_happened_is_on_the_first_screen(self):
        md = render.render_markdown(self.grouped())
        self.assertIn("**1 stories**", md)
        self.assertIn("1 epics", md)
        self.assertIn("2 source rows", md)

    def test_every_source_row_is_still_named_under_its_story(self):
        """A reviewer checks the inventory against the client's own document line by line,
        which is only possible while every line is still on the page."""
        md = render.render_markdown(self.grouped())
        self.assertIn("#### T1 — Search leads", md)
        self.assertIn("row 182", md)

    def test_a_source_row_carries_its_text_and_not_only_a_reference_to_it(self):
        """The defect this exists to stop: 912 rows rendered as `id + location` with the
        client's own sentence — the thing an estimator sizes against — left in the source."""
        md = render.render_markdown(self.grouped())
        self.assertIn("Search the leads list by name.", md)
        self.assertIn("Filter leads by status.", md)

    def test_surfaces_say_which_roles_are_on_the_work(self):
        self.assertIn("only these roles are billed", render.render_markdown(self.grouped()))

    def test_implicit_scope_shows_its_reason_where_a_quote_would_be(self):
        inv = self.grouped()
        item = feature("I1", "Migrate ten years of bookings")
        del item["citations"]
        item["rationale"] = "The workbook says the bookings already exist."
        inv["implicit_scope"] = [item]
        md = render.render_markdown(inv)
        self.assertIn("no quote behind it", md)
        self.assertIn("bookings already exist", md)


class Pages(unittest.TestCase):
    """One file per epic. Inline row text takes a 365-story inventory past 2 MB, and a
    document nobody can open is not a readable one."""

    def grouped(self, n=2):
        features = [feature(f"F{i}", f"Story {i}", epic_id=f"E{i % n + 1}") for i in range(1, 5)]
        epics = [{"id": f"E{i}", "name": f"Epic {i}", "origin": "source"} for i in range(1, n + 1)]
        return inventory(features=features, epics=epics)

    def test_an_ungrouped_inventory_stays_one_file(self):
        """check_grouping enforces all-or-nothing grouping, so an inventory with no epics has
        nothing to split on — a directory of one file would be worse than the file."""
        pages, _ = render.plan_pages(inventory())
        self.assertEqual([p["path"] for p in pages], ["feature-inventory.md"])
        self.assertIn("Users must be able to log in", render.render_pages(inventory())["feature-inventory.md"])

    def test_a_grouped_inventory_splits_one_file_per_epic(self):
        out = render.render_pages(self.grouped())
        self.assertEqual(sorted(out), ["feature-inventory.md",
                                       "inventory/e1-epic-1.md", "inventory/e2-epic-2.md"])

    def test_the_index_links_to_every_page_and_holds_no_stories(self):
        out = render.render_pages(self.grouped())
        index = out["feature-inventory.md"]
        self.assertIn("## Contents", index)
        for path in out:
            if path != "feature-inventory.md":
                self.assertIn(f"]({path})", index)
        self.assertNotIn("### F1 —", index)

    def test_every_epic_page_links_back_to_the_index(self):
        out = render.render_pages(self.grouped())
        for path, text in out.items():
            if path != "feature-inventory.md":
                self.assertIn("(../feature-inventory.md)", text)

    def test_a_story_with_no_epic_still_gets_a_page(self):
        inv = self.grouped()
        inv["features"].append(feature("F9", "Orphan"))
        out = render.render_pages(inv)
        self.assertIn("inventory/unassigned.md", out)
        self.assertIn("### F9 — Orphan", out["inventory/unassigned.md"])

    def test_an_empty_epic_gets_no_page(self):
        inv = self.grouped()
        inv["epics"].append({"id": "E9", "name": "Nothing here", "origin": "source"})
        self.assertNotIn("inventory/e9-nothing-here.md", render.render_pages(inv))

    def test_every_story_carries_an_explicit_anchor(self):
        """Heading-derived anchors change whenever a story is renamed. A cross-file link has
        to survive a rename, so the id is written out."""
        out = render.render_pages(self.grouped())
        self.assertIn('<a id="F1"></a>', out["inventory/e2-epic-2.md"])


class References(unittest.TestCase):
    """Nothing renders as a bare id."""

    def linked(self):
        f1 = feature("F1", "Sign in", epic_id="E1")
        f2 = feature("F2", "Reset password", epic_id="E2", depends_on=[
            {"feature_id": "F1", "inferred": False,
             "evidence": "Password reset is only offered to a signed-in user."},
            {"feature_id": "F9", "inferred": True, "why": "Deduced from the audit requirement."}])
        f9 = feature("F9", "Audit log", epic_id="E2")
        return inventory(features=[f1, f2, f9],
                         epics=[{"id": "E1", "name": "Auth", "origin": "source"},
                                {"id": "E2", "name": "Account", "origin": "source"}])

    def test_a_dependency_renders_its_evidence_not_a_bare_id(self):
        """307 dependencies in a real run carried the sentence that stated them, and both
        renderers dropped it on the way out."""
        out = render.render_pages(self.linked())
        page = out["inventory/e2-account.md"]
        self.assertIn("Password reset is only offered to a signed-in user.", page)
        self.assertIn("F1 — Sign in", page)

    def test_an_inferred_dependency_renders_the_reason_it_was_deduced(self):
        page = render.render_pages(self.linked())["inventory/e2-account.md"]
        self.assertIn("Deduced from the audit requirement.", page)
        self.assertIn("*(inferred)*", page)

    def test_a_cross_file_dependency_link_resolves_to_a_real_anchor(self):
        out = render.render_pages(self.linked())
        page = out["inventory/e2-account.md"]
        self.assertIn("(e1-auth.md#F1)", page)   # sibling pages, same directory
        self.assertIn('<a id="F1"></a>', out["inventory/e1-auth.md"])

    def test_a_dependency_within_the_same_page_links_inside_it(self):
        page = render.render_pages(self.linked())["inventory/e2-account.md"]
        self.assertIn("(e2-account.md#F9)", page)

    def test_the_csv_dependency_column_carries_the_name_and_the_reason(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "f.csv"
            render.render_csv(self.linked(), target)
            rows = {r["id"]: r for r in read_csv(target)}
        self.assertIn("F1 (Sign in)", rows["F2"]["depends_on"])
        self.assertIn("only offered to a signed-in user", rows["F2"]["depends_on"])
        self.assertIn("[inferred]", rows["F2"]["depends_on"])


class SourceLinks(unittest.TestCase):
    """A location is a place in a document. It should be one click, not a search."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        (self.dir / "normalized").mkdir()
        (self.dir / "normalized" / "S1-doc.md").write_text(NORMALIZED, encoding="utf-8")
        f = feature("F1", "Domain access", epic_id="E1", citations=[
            {"source_id": "S1", "location": "sheet 'Operators - Feature list' row 9",
             "quote": ROW9}],
            tasks=[task("F1-T9", "Enforce domain access",
                        "sheet 'Operators - Feature list' row 9", ROW9),
                   task("F1-T10", "Enforce domain access",
                        "sheet 'Operators - Feature list' row 10", ROW10)])
        self.inv = inventory(features=[f], epics=[{"id": "E1", "name": "Sys", "origin": "source"}],
                             sources=[{"id": "S1", "path": "docs/book.xlsx", "doc_type": "backlog",
                                       "language": "en", "converter": "openpyxl"}])
        self.links = render.Links(self.dir / "normalized", self.inv)

    def tearDown(self):
        self.tmp.cleanup()

    def test_a_link_lands_on_the_line_the_row_is_actually_on(self):
        """A link to the wrong line is worse than no link: it looks checked and is not."""
        lines = NORMALIZED.splitlines()
        for row, cit in ((9, {"source_id": "S1", "location": "sheet 'Operators - Feature list' row 9"}),
                         (10, {"source_id": "S1", "location": "sheet 'Operators - Feature list' row 10"})):
            n = self.links.line_of(cit)
            self.assertIsNotNone(n, row)
            self.assertRegex(lines[n - 1], rf"^\|\s*{row}\s*\|")

    def test_the_href_is_relative_to_the_page_it_is_written_into(self):
        cit = {"source_id": "S1", "location": "sheet 'Operators - Feature list' row 9"}
        self.assertEqual(self.links.href(cit, "feature-inventory.md"), "normalized/S1-doc.md#L7")
        self.assertEqual(self.links.href(cit, "inventory/e1-sys.md"), "../normalized/S1-doc.md#L7")

    def test_every_link_in_a_rendered_page_resolves(self):
        out = render.render_pages(self.inv, self.links)
        for path, text in out.items():
            target = self.dir / path
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(text, encoding="utf-8")
        found = 0
        for path, text in out.items():
            for href in re.findall(r"\]\(([^)]+)\)", text):
                target, _, frag = href.partition("#")
                resolved = (self.dir / path).parent / target
                self.assertTrue(resolved.exists(), f"{path} -> {href}")
                if frag.startswith("L"):
                    found += 1
                    self.assertLessEqual(int(frag[1:]),
                                         len(resolved.read_text().splitlines()))
        self.assertGreater(found, 0, "no source links were emitted at all")

    def test_without_normalized_sources_it_renders_plain_text_rather_than_failing(self):
        md = render.render_markdown(self.inv)
        self.assertIn("sheet 'Operators - Feature list' row 9", md)
        self.assertNotIn("](normalized", md)

    def test_a_story_citation_is_suppressed_only_when_a_task_already_carries_it_whole(self):
        """The story and its row quote the same cell. Printing both is noise — but where the
        story's quote is the longer of the two, dropping it would lose the text."""
        page = render.render_pages(self.inv, self.links)["inventory/e1-sys.md"]
        self.assertEqual(page.count(ROW9), 1, "the same passage printed twice")

        trimmed = json.loads(json.dumps(self.inv))
        trimmed["features"][0]["tasks"][0]["citations"][0]["quote"] = "Enforce domain access"
        page = render.render_pages(trimmed, self.links)["inventory/e1-sys.md"]
        self.assertIn(ROW9, page, "the fuller quote was dropped as a duplicate of a shorter one")


class Tasks(unittest.TestCase):
    """`F1-TO2` used to appear in exactly one cell of the CSV and match no row in it."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        f = feature("F1", "Domain access", epic_id="E1",
                    tasks=[task("F1-T9", "Enforce access",
                                "sheet 'Operators - Feature list' row 9", ROW9),
                           task("F1-T10", "Enforce access",
                                "sheet 'Operators - Feature list' row 10", ROW10)])
        self.inv = inventory(features=[f], epics=[{"id": "E1", "name": "Sys", "origin": "source"}])

    def tearDown(self):
        self.tmp.cleanup()

    def rows(self):
        target = self.dir / "t.csv"
        render.render_tasks_csv(self.inv, target)
        return read_csv(target)

    def test_one_row_per_task_carrying_its_text(self):
        rows = self.rows()
        self.assertEqual([r["task_id"] for r in rows], ["F1-T9", "F1-T10"])
        self.assertEqual(rows[0]["text"], ROW9)

    def test_the_join_resolves_in_both_directions(self):
        stories = self.dir / "s.csv"
        render.render_csv(self.inv, stories)
        story = read_csv(stories)[0]
        tasks = {r["task_id"]: r for r in self.rows()}
        for tid in story["task_ids"].split("; "):
            self.assertIn(tid, tasks, "a story names a task with no row of its own")
            self.assertEqual(tasks[tid]["feature_id"], story["id"])
        self.assertEqual(story["task_count"], "2")

    def test_the_source_is_broken_out_into_columns_that_sort(self):
        row = self.rows()[0]
        self.assertEqual((row["sheet"], row["row"]), ("Operators - Feature list", "9"))
        self.assertEqual(row["source_id"], "S1")

    def test_a_story_with_no_tasks_produces_no_task_rows(self):
        self.inv["features"][0].pop("tasks")
        self.assertEqual(self.rows(), [])


@unittest.skipUnless(HAS_XLSX, "openpyxl not importable")
class Workbook(unittest.TestCase):
    """Two tabs, and a reference you follow rather than resolve by hand."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        f1 = feature("F1", "Domain access", epic_id="E1",
                     tasks=[task("F1-T9", "Enforce access", "sheet 'S' row 9", ROW9)])
        f2 = feature("F2", "Audit", epic_id="E1",
                     tasks=[task("F2-T1", "Write entries", "sheet 'S' row 12", "Record writes.")])
        self.inv = inventory(features=[f1, f2],
                             epics=[{"id": "E1", "name": "Sys", "origin": "source"}])
        self.target = self.dir / "book.xlsx"
        self.assertTrue(render.render_xlsx(self.inv, self.target))

    def tearDown(self):
        self.tmp.cleanup()

    def book(self):
        from openpyxl import load_workbook
        return load_workbook(self.target)

    def test_two_sheets_with_a_row_for_every_story_and_every_task(self):
        wb = self.book()
        self.assertEqual(wb.sheetnames, ["Stories", "Tasks"])
        self.assertEqual(wb["Stories"].max_row - 1, 2)
        self.assertEqual(wb["Tasks"].max_row - 1, 2)

    def test_a_story_clicks_through_to_its_own_first_task(self):
        wb = self.book()
        st, tk = wb["Stories"], wb["Tasks"]
        header = [c.value for c in st[1]]
        cell = st.cell(row=3, column=header.index("task_ids") + 1)   # F2
        self.assertIsNotNone(cell.hyperlink, "task_ids is not a link")
        row = int(cell.hyperlink.location.split("!A")[1])
        self.assertEqual(tk.cell(row=row, column=1).value, "F2-T1")

    def test_a_task_clicks_back_to_the_story_it_belongs_to(self):
        wb = self.book()
        st, tk = wb["Stories"], wb["Tasks"]
        header = [c.value for c in tk[1]]
        cell = tk.cell(row=3, column=header.index("feature_id") + 1)
        row = int(cell.hyperlink.location.split("!A")[1])
        self.assertEqual(st.cell(row=row, column=1).value, cell.value)

    def test_the_task_text_is_in_the_workbook_not_only_an_id(self):
        wb = self.book()
        header = [c.value for c in wb["Tasks"][1]]
        self.assertEqual(wb["Tasks"].cell(row=2, column=header.index("text") + 1).value, ROW9)

    def test_the_header_is_frozen_and_filterable_so_it_can_be_worked_in(self):
        st = self.book()["Stories"]
        self.assertEqual(st.freeze_panes, "A2")
        self.assertTrue(st.auto_filter.ref)

    def test_a_source_link_is_a_hyperlink_when_the_sources_are_available(self):
        (self.dir / "normalized").mkdir()
        (self.dir / "normalized" / "S1-doc.md").write_text(NORMALIZED, encoding="utf-8")
        inv = json.loads(json.dumps(self.inv))
        inv["features"][0]["tasks"][0]["citations"][0]["location"] = \
            "sheet 'Operators - Feature list' row 9"
        links = render.Links(self.dir / "normalized", inv)
        render.render_xlsx(inv, self.target, links)
        tk = self.book()["Tasks"]
        header = [c.value for c in tk[1]]
        cell = tk.cell(row=2, column=header.index("link") + 1)
        self.assertTrue(cell.value.endswith("#L7"), cell.value)
        self.assertIsNotNone(cell.hyperlink)


class WithoutOpenpyxl(unittest.TestCase):
    def test_the_render_still_succeeds_and_says_the_workbook_was_skipped(self):
        """`uv run` provisions openpyxl; bare python3 may not have it. Losing the workbook is
        a degraded render, not a failed one — the same bargain convert-input.py already makes."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "feature-inventory.json"
            path.write_text(json.dumps(inventory()), encoding="utf-8")
            saved = sys.modules.get("openpyxl")
            sys.modules["openpyxl"] = None
            try:
                sys.argv = ["render-inventory.py", str(path), "--formats", "csv,xlsx"]
                out = io.StringIO()
                with contextlib.redirect_stdout(out):
                    self.assertEqual(render.main(), 0)
                result = json.loads(out.getvalue())
            finally:
                if saved is None:
                    sys.modules.pop("openpyxl", None)
                else:
                    sys.modules["openpyxl"] = saved
            self.assertFalse((Path(tmp) / "feature-inventory.xlsx").exists())
            self.assertTrue((Path(tmp) / "feature-inventory.csv").exists())
            self.assertTrue(any("xlsx skipped" in n for n in result["notes"]))


if __name__ == "__main__":
    unittest.main()
