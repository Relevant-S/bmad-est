#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Render a Feature Inventory into the human and sales views.

feature-inventory.json is the source of truth; the .md, .csv and .xlsx are projections
regenerated on demand rather than maintained. Features render as sections rather than
table rows because verbatim citations carry pipes, newlines and non-Latin text that a
markdown table destroys — and destroying the quote destroys the traceability.

**Every reference resolves.** The projections used to emit pointers: task ids that appeared
in no row of the file that named them, dependencies as bare feature ids, locations as prose
a reader had to go and find by hand, and `evidence` — the sentence stating a dependency,
already sitting in the JSON — dropped on the way out. A reference nobody can follow is worse
than no reference, because it reads as though the work of checking has been done.

These views carry scope, not cost. The classification a story is priced on lives in
classification.json and renders through est-estimate, so a reviewer reading this is reading
what the client asked for with nothing about effort mixed into it.
"""

import argparse
import csv
import importlib.util
import json
import posixpath
import re
import sys
from pathlib import Path

SCOPE_MARK = {"outside_agreed_scope": " · outside agreed scope",
              "no_agreed_scope_defined": " · no agreed scope defined"}

INDEX_NAME = "feature-inventory.md"
PAGES_DIR = "inventory"
TASKS_CSV = "feature-inventory.tasks.csv"

STORY_COLUMNS = ["id", "name", "description", "epic_id", "epic_name", "surfaces",
                 "commitment", "scope_status", "origin", "task_count", "task_ids",
                 "depends_on", "open_questions", "sources", "locations", "quotes"]

TASK_COLUMNS = ["task_id", "feature_id", "feature_name", "epic_id", "name", "text",
                "source_id", "source_path", "sheet", "row", "location", "link"]


# --- locating a citation in the converted source -------------------------------------------

def _load_checker():
    """`inventory-check.py` owns what a location means. Import it rather than re-deriving it,
    the way classification-merge.py imports the diff's matcher: two definitions of where a
    citation points would drift, and the links would drift with them."""
    path = Path(__file__).resolve().parent / "inventory-check.py"
    spec = importlib.util.spec_from_file_location("inventory_check", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class Links:
    """Turns a citation into a relative link that lands on the row it cites.

    Without a `--normalized` directory every method returns None and the renderers fall back
    to plain text, so the projections still render when the converted sources are gone.
    """

    def __init__(self, normalized_dir=None, inv=None):
        self.anchors, self.files = {}, {}
        if not normalized_dir:
            return
        check = _load_checker()
        self.loc_sheet_row = check.LOC_SHEET_ROW
        self.loc_page, self.loc_slide = check.LOC_PAGE, check.LOC_SLIDE
        root = Path(normalized_dir)
        for source in (inv or {}).get("sources", []):
            sid = source.get("id")
            found = sorted(root.glob(f"{sid}-*.md"))
            if not found:
                continue
            self.files[sid] = f"{root.name}/{found[0].name}"
            self.anchors[sid] = check.index_anchors(
                found[0].read_text(encoding="utf-8", errors="replace"))

    def line_of(self, cit):
        """The line in the converted source this citation points at, if it can be resolved."""
        sid = cit.get("source_id")
        if sid not in self.anchors:
            return None
        location, anchors = cit.get("location") or "", self.anchors[sid]
        match = self.loc_sheet_row.search(location)
        if match:
            return anchors["sheets"].get(match.group(1), {}).get(int(match.group(2)))
        for pattern, key in ((self.loc_slide, "slides"), (self.loc_page, "pages")):
            match = pattern.search(location)
            if match:
                return anchors[key].get(int(match.group(1)))
        return None

    def href(self, cit, from_page):
        """A link to the cited row, relative to the page it is being written into."""
        sid = cit.get("source_id")
        if sid not in self.files:
            return None
        target = _relative(self.files[sid], from_page)
        line = self.line_of(cit)
        return f"{target}#L{line}" if line else target


def _relative(target, from_page):
    """`target` and `from_page` are both workspace-relative POSIX paths."""
    base = posixpath.dirname(from_page)
    return posixpath.relpath(target, base) if base else target


def _slug(text, fallback="untitled"):
    slug = re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")
    return slug[:60] or fallback


def _flat(text):
    return re.sub(r"\s+", " ", (text or "")).strip().lower()


def _sheet_row(cit, links):
    """The sheet and row a location names, when it names one."""
    pattern = getattr(links, "loc_sheet_row", None) or re.compile(
        r"sheet\s*['\"]?(.+?)['\"]?\s*[, ]\s*row\s*(\d+)", re.I)
    match = pattern.search(cit.get("location") or "")
    return (match.group(1), match.group(2)) if match else ("", "")


# --- the page layout ------------------------------------------------------------------------

def plan_pages(inv):
    """Which file each story is rendered into.

    One file per epic, because inline row text takes a 365-story inventory past 2 MB and a
    document nobody opens is not a readable one. An inventory with no epics has nothing to
    split on — `check_grouping` enforces all-or-nothing grouping — so it stays a single file
    rather than becoming a directory of one.
    """
    features = list(inv.get("features") or [])
    implicit = list(inv.get("implicit_scope") or [])
    epics = list(inv.get("epics") or [])
    if not epics or not any(f.get("epic_id") for f in features):
        return [{"path": INDEX_NAME, "title": None, "epic": None,
                 "features": features + implicit}], {}

    pages, by_feature = [], {}
    for epic in epics:
        eid = epic.get("id")
        members = [f for f in features if f.get("epic_id") == eid]
        pages.append({"path": f"{PAGES_DIR}/{_slug(str(eid))}-{_slug(epic.get('name'))}.md",
                      "title": f"{eid} — {epic.get('name')}", "epic": epic, "features": members})
    loose = [f for f in features if not f.get("epic_id")]
    if loose:
        pages.append({"path": f"{PAGES_DIR}/unassigned.md", "title": "Stories with no epic",
                      "epic": None, "features": loose})
    if implicit:
        pages.append({"path": f"{PAGES_DIR}/implicit.md", "title": "Implicit scope",
                      "epic": None, "features": implicit})

    pages = [p for p in pages if p["features"]]
    for page in pages:
        for f in page["features"]:
            by_feature[f.get("id")] = page["path"]
    return pages, by_feature


# --- markdown -------------------------------------------------------------------------------

def _citation_line(cit, src_by_id, links, page, indent=""):
    sid = cit.get("source_id")
    doc = Path(src_by_id.get(sid, {}).get("path", sid or "")).name or sid
    label = f"{sid} · {doc} — {cit.get('location')}"
    href = links.href(cit, page) if links else None
    head = f"{indent}- [{label}]({href})" if href else f"{indent}- {label}"
    out = [head, f"{indent}  > {(cit.get('quote') or '').strip()}"]
    if cit.get("quote_original"):
        out.append(f"{indent}  > *({cit.get('quote_language', 'original')})* "
                   f"{cit['quote_original'].strip()}")
    return out


def story_block(f, inv, links, page, by_feature, src_by_id, names):
    out = [f"<a id=\"{f.get('id')}\"></a>",
           f"### {f.get('id')} — {f.get('name')}{SCOPE_MARK.get(f.get('scope_status'), '')}",
           "", f.get("description", ""), ""]

    if f.get("surfaces"):
        out.append(f"- **Surfaces:** {', '.join(f['surfaces'])} — only these roles are billed to it.")
    out.append(f"- **Commitment:** {f.get('commitment')}")
    if f.get("rationale"):
        out.append(f"- **Why it is here, with no quote behind it:** {f['rationale']}")
    if f.get("scope_status") == "outside_agreed_scope":
        out.append("- **Outside the agreed scope** — estimated separately, not folded into "
                   "the agreed number.")

    deps = f.get("depends_on") or []
    if deps:
        # The evidence was always in the JSON and both renderers threw it away, printing a bare
        # id. A dependency you cannot see the reason for is a second lookup, not information.
        out += ["", "**Depends on:**", ""]
        for dep in deps:
            target = dep.get("feature_id")
            label = f"{target} — {names[target]}" if target in names else str(target)
            href = by_feature.get(target)
            link = f"[{label}]({_relative(href, page)}#{target})" if href else label
            out.append(f"- {link}{' *(inferred)*' if dep.get('inferred') else ''}")
            reason = (dep.get("evidence") or "").strip() or (dep.get("why") or "").strip()
            if reason:
                out.append(f"  > {reason}")
        out.append("")

    rows = f.get("tasks") or []
    shown = {}
    if rows:
        out += ["", f"**Source rows ({len(rows)}):**", ""]
        for task in rows:
            out += [f"<a id=\"{task.get('id')}\"></a>",
                    f"#### {task.get('id')} — {task.get('name')}", ""]
            for cit in task.get("citations") or []:
                key = (cit.get("source_id"), cit.get("location"))
                shown.setdefault(key, "")
                shown[key] += " " + _flat(cit.get("quote"))
                out += _citation_line(cit, src_by_id, links, page)
            out.append("")

    # A story's own citations, minus the rows already written out above it — printing the same
    # passage twice is how a document that carries everything becomes unreadable. Suppressed
    # only when the task's quote already contains the story's: where the two quote the same
    # cell to different lengths, dropping the longer one would lose the text this exists to keep.
    extra = [c for c in f.get("citations", [])
             if _flat(c.get("quote")) not in shown.get((c.get("source_id"), c.get("location")), "\x00")]
    if extra:
        out += ["", "**Cited:**" if rows else "**Source:**", ""]
        for cit in extra:
            out += _citation_line(cit, src_by_id, links, page)
        out.append("")

    if f.get("open_questions"):
        out += ["", "**Open questions:**", ""]
        out += [f"- ❓ {q}" for q in f["open_questions"]]
        out.append("")
    return out


def render_index(inv, pages, links):
    src_by_id = {s["id"]: s for s in inv.get("sources", [])}
    features = list(inv.get("features") or [])
    implicit = list(inv.get("implicit_scope") or [])
    epics = inv.get("epics") or []
    split = len(pages) > 1 or pages[0]["path"] != INDEX_NAME

    out = [f"# Feature Inventory — {inv.get('project', 'untitled')}", ""]
    out.append(f"Generated {inv.get('generated', '')} · schema {inv.get('schema_version', '')} · "
               f"granularity: {inv.get('granularity', 'unspecified')}")

    tasks = sum(len(f.get("tasks") or []) for f in features)
    # Story count, epic count and source-row count on one line, because the ratio between them
    # is the fastest way to see whether the workbook was grouped or transliterated.
    counts = [f"**{len(features)} stories**"]
    if epics:
        counts.append(f"{len(epics)} epics")
    if tasks:
        counts.append(f"{tasks} source rows")
    if implicit:
        counts.append(f"{len(implicit)} implicit")
    counts += [f"{len(inv.get('not_scope', []))} passages excluded from scope",
               f"{len(inv.get('conflicts', []))} conflicts"]
    out += ["", " · ".join(counts), ""]

    if split:
        out += ["## Contents", ""]
        for page in pages:
            n = len(page["features"])
            out.append(f"- [{page['title']}]({page['path']}) · {n} "
                       f"{'story' if n == 1 else 'stories'}")
        out.append("")

    if epics:
        out += ["## Epics", "", "| id | epic | stories | from |", "| --- | --- | ---: | --- |"]
        page_of = {p["epic"]["id"]: p["path"] for p in pages if p.get("epic")}
        for epic in epics:
            eid = epic.get("id")
            n = sum(1 for f in features if f.get("epic_id") == eid)
            origin = epic.get("origin", "")
            if origin == "synthesised" and epic.get("why"):
                origin += f" — {epic['why']}"
            name = epic.get("name")
            if eid in page_of:
                name = f"[{name}]({page_of[eid]})"
            out.append(f"| {eid} | {name} | {n} | {origin} |")
        out.append("")

    out += ["## Sources", "", "| id | document | type | language | converter | coverage note |",
            "| --- | --- | --- | --- | --- | --- |"]
    for s in inv.get("sources", []):
        path = s.get("path")
        href = links.files.get(s.get("id")) if links else None
        shown = f"[`{path}`]({href})" if href else f"`{path}`"
        out.append(f"| {s.get('id')} | {shown} | {s.get('doc_type')} | {s.get('language')} "
                   f"| {s.get('converter', '')} | {s.get('coverage_note', '')} |")
    out.append("")

    if inv.get("not_scope"):
        out += ["## Not treated as scope", "",
                "Source content that did not become a feature, and why.", "",
                "| source | location | reason | text |", "| --- | --- | --- | --- |"]
        for ns in inv["not_scope"]:
            quote = (ns.get("quote", "") or "").replace("|", "\\|").replace("\n", " ")
            quote = quote[:200] + ("…" if len(quote) > 200 else "")
            out.append(f"| {ns.get('source_id')} | {ns.get('location')} | {ns.get('reason')} "
                       f"| {quote} |")
        out.append("")

    if inv.get("conflicts"):
        out += ["## Conflicts between sources", ""]
        for c in inv["conflicts"]:
            out.append(f"- **{c.get('topic')}** ({', '.join(c.get('source_ids', []))}) "
                       f"— {c.get('detail')}")
        out.append("")

    if inv.get("assumptions"):
        out += ["## Assumptions", ""] + [f"- {a}" for a in inv["assumptions"]] + [""]
    return out


def render_pages(inv, links=None, pages=None, by_feature=None):
    """Every markdown page this inventory renders to, keyed by workspace-relative path."""
    links = links or Links()
    if pages is None:
        pages, by_feature = plan_pages(inv)
    src_by_id = {s["id"]: s for s in inv.get("sources", [])}
    names = {f.get("id"): f.get("name")
             for f in list(inv.get("features") or []) + list(inv.get("implicit_scope") or [])}

    out = {}
    index = render_index(inv, pages, links)
    single = len(pages) == 1 and pages[0]["path"] == INDEX_NAME
    if single:
        index += ["## Stories", ""]
        for f in pages[0]["features"]:
            index += story_block(f, inv, links, INDEX_NAME, by_feature, src_by_id, names)
        out[INDEX_NAME] = "\n".join(index)
        return out

    out[INDEX_NAME] = "\n".join(index)
    for page in pages:
        body = [f"# {page['title']}", "",
                f"[← Feature Inventory — {inv.get('project', 'untitled')}]"
                f"({_relative(INDEX_NAME, page['path'])})", ""]
        if page.get("epic") and page["epic"].get("origin") == "synthesised" and page["epic"].get("why"):
            body += [f"*Synthesised epic — {page['epic']['why']}*", ""]
        for f in page["features"]:
            body += story_block(f, inv, links, page["path"], by_feature, src_by_id, names)
        out[page["path"]] = "\n".join(body)
    return out


def render_markdown(inv, links=None):
    """Every page as one string — what the single-file layout writes verbatim."""
    return "\n".join(render_pages(inv, links).values())


# --- tabular --------------------------------------------------------------------------------

def story_rows(inv, links=None):
    links = links or Links()
    src_by_id = {s["id"]: s for s in inv.get("sources", [])}
    epic_names = {e.get("id"): e.get("name") for e in inv.get("epics") or []}
    names = {f.get("id"): f.get("name")
             for f in list(inv.get("features") or []) + list(inv.get("implicit_scope") or [])}

    rows = []
    for f in list(inv.get("features", [])) + list(inv.get("implicit_scope") or []):
        citations = f.get("citations", []) or []
        deps = []
        for d in f.get("depends_on", []) or []:
            target = d.get("feature_id")
            label = f"{target} ({names[target]})" if target in names else str(target)
            reason = (d.get("evidence") or "").strip() or (d.get("why") or "").strip()
            deps.append(f"{label}{' [inferred]' if d.get('inferred') else ''}"
                        f"{': ' + reason if reason else ''}")
        rows.append({
            "id": f.get("id"),
            "name": f.get("name"),
            "description": f.get("description"),
            "epic_id": f.get("epic_id") or "",
            "epic_name": epic_names.get(f.get("epic_id"), ""),
            "surfaces": "; ".join(f.get("surfaces") or []),
            "commitment": f.get("commitment"),
            "scope_status": f.get("scope_status") or "in_agreed_scope",
            "origin": f.get("origin") or ("implicit" if f.get("rationale") else "extracted"),
            "task_count": len(f.get("tasks") or []),
            "task_ids": "; ".join(t.get("id", "") for t in f.get("tasks") or []),
            "depends_on": "\n".join(deps),
            "open_questions": "\n".join(f.get("open_questions") or []),
            "sources": "; ".join(dict.fromkeys(
                Path(src_by_id.get(c.get("source_id"), {}).get("path", c.get("source_id") or "")).name
                for c in citations)),
            "locations": "; ".join(c.get("location", "") for c in citations),
            # Every quote, not the first. `primary_quote` silently dropped the rest, so a story
            # assembled from three rows published one of them.
            "quotes": "\n\n".join(f"[{c.get('location')}] {(c.get('quote') or '').strip()}"
                                  for c in citations),
        })
    return rows


def task_rows(inv, links=None, from_page=""):
    links = links or Links()
    src_by_id = {s["id"]: s for s in inv.get("sources", [])}
    rows = []
    for f in list(inv.get("features", [])) + list(inv.get("implicit_scope") or []):
        for task in f.get("tasks") or []:
            citations = task.get("citations") or []
            first = citations[0] if citations else {}
            sheet, row = _sheet_row(first, links)
            rows.append({
                "task_id": task.get("id"),
                "feature_id": f.get("id"),
                "feature_name": f.get("name"),
                "epic_id": f.get("epic_id") or "",
                "name": task.get("name"),
                "text": "\n\n".join((c.get("quote") or "").strip() for c in citations),
                "source_id": first.get("source_id", ""),
                "source_path": src_by_id.get(first.get("source_id"), {}).get("path", ""),
                "sheet": sheet,
                "row": row,
                "location": "; ".join(c.get("location", "") for c in citations),
                "link": links.href(first, from_page) or "" if citations else "",
            })
    return rows


def render_csv(inv, target, links=None):
    with open(target, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=STORY_COLUMNS)
        writer.writeheader()
        writer.writerows(story_rows(inv, links))


def render_tasks_csv(inv, target, links=None):
    with open(target, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=TASK_COLUMNS)
        writer.writeheader()
        writer.writerows(task_rows(inv, links))


WIDE = {"description": 70, "quotes": 90, "text": 90, "depends_on": 50, "open_questions": 50,
        "name": 42, "feature_name": 42, "link": 46, "sheet": 26, "source_path": 30}


WRAP = ("description", "quotes", "text", "depends_on", "open_questions", "why")


def write_workbook(sheets, target, links=()):
    """Write one workbook from a list of (title, columns, rows), with the cross-links.

    Shared with est-estimate, which puts the priced estimate on the same two tabs. Two writers
    would drift, and the point of the estimate workbook is that it *is* the inventory workbook
    with more columns.

    `links` is a list of (from_sheet, from_column, to_sheet, key_column, index) tuples, where
    `index` maps a key to the 1-based row it should jump to. Returns False if openpyxl is
    absent, so the caller can degrade to CSV rather than fail.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.hyperlink import Hyperlink
    except ImportError:
        return False

    wb = Workbook()
    made = {}
    for n, (title, columns, rows) in enumerate(sheets):
        sheet = wb.active if n == 0 else wb.create_sheet(title)
        sheet.title = title
        made[title] = (sheet, columns, rows)
        sheet.append(list(columns))
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([row.get(c, "") for c in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = (f"A1:{get_column_letter(len(columns))}"
                                 f"{max(len(rows) + 1, 1)}")
        for i, name in enumerate(columns, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = WIDE.get(name, 18)
        wrap = Alignment(wrap_text=True, vertical="top")
        for name in WRAP:
            if name not in columns:
                continue
            letter = get_column_letter(columns.index(name) + 1)
            for cell in sheet[letter][1:]:
                cell.alignment = wrap

    for from_sheet, from_column, to_sheet, key_column, index in links:
        if from_sheet not in made or from_column not in made[from_sheet][1]:
            continue
        sheet, columns, rows = made[from_sheet]
        col = columns.index(from_column) + 1
        for i, row in enumerate(rows, start=2):
            cell = sheet.cell(row=i, column=col)
            if to_sheet is None:                      # an external link, in the cell's own value
                if row.get(from_column):
                    cell.hyperlink = row[from_column]
                    cell.style = "Hyperlink"
                continue
            at = index.get(row.get(key_column))
            if at:
                cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"{to_sheet}!A{at}")
                cell.style = "Hyperlink"

    wb.save(target)
    return True


def workbook_links(stories, tasks, story_key="id", task_parent="feature_id",
                   story_link_col="task_ids"):
    """The two-way index between a Stories sheet and a Tasks sheet."""
    first_task = {}
    for i, row in enumerate(tasks, start=2):
        first_task.setdefault(row.get(task_parent), i)
    story_row = {row.get(story_key): i for i, row in enumerate(stories, start=2)}
    return [("Stories", story_link_col, "Tasks", story_key, first_task),
            ("Tasks", task_parent, "Stories", task_parent, story_row),
            ("Tasks", "link", None, None, {})]


def render_xlsx(inv, target, links=None):
    """Two sheets with hyperlinks in both directions, and out to the converted source.

    A CSV can only ever hand you an id to go and look up: `F1-TO2` appeared in exactly one cell
    of the old file and matched no row in it. Here the story's task cell clicks through to its
    rows and each row clicks back, so a reference is something you follow rather than resolve.
    """
    stories, tasks = story_rows(inv, links), task_rows(inv, links)
    return write_workbook(
        [("Stories", STORY_COLUMNS, stories), ("Tasks", TASK_COLUMNS, tasks)],
        target, workbook_links(stories, tasks))


def main():
    ap = argparse.ArgumentParser(
        description="Render feature-inventory.json into its markdown, CSV and workbook views.",
        epilog="Exit codes: 0 rendered, 2 unreadable input.",
    )
    ap.add_argument("inventory", help="path to feature-inventory.json")
    ap.add_argument("--out-dir", help="directory for the rendered files (default: alongside the inventory)")
    ap.add_argument("--normalized",
                    help="the workspace's normalized/ directory. Without it citations render as "
                         "plain text instead of links into the converted source.")
    ap.add_argument("--formats", default="md,csv,xlsx",
                    help="comma-separated subset of md,csv,xlsx. Members this script does not "
                         "produce are ignored, so the shared est_output_formats value can be "
                         "passed through unchanged.")
    args = ap.parse_args()

    path = Path(args.inventory)
    try:
        inv = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        print(json.dumps({"ok": False, "error": f"cannot read {path}: {exc}"}))
        return 2

    out_dir = Path(args.out_dir) if args.out_dir else path.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    formats = {f.strip() for f in args.formats.split(",") if f.strip()}
    links = Links(args.normalized, inv) if args.normalized else Links()
    written, notes = [], []

    if "md" in formats:
        pages, by_feature = plan_pages(inv)
        for rel, text in render_pages(inv, links, pages, by_feature).items():
            target = out_dir / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(text + "\n", encoding="utf-8")
            written.append(str(target))
    if "csv" in formats:
        render_csv(inv, out_dir / "feature-inventory.csv", links)
        render_tasks_csv(inv, out_dir / TASKS_CSV, links)
        written += [str(out_dir / "feature-inventory.csv"), str(out_dir / TASKS_CSV)]
    if "xlsx" in formats:
        target = out_dir / "feature-inventory.xlsx"
        if render_xlsx(inv, target, links):
            written.append(str(target))
        else:
            notes.append("xlsx skipped: openpyxl is not importable. Run this under `uv run`, "
                         "which provisions it, or drop xlsx from est_output_formats.")
    if not args.normalized:
        notes.append("no --normalized: citations rendered as plain text, without links into "
                     "the converted sources.")

    result = {"ok": True, "written": written, "features": len(inv.get("features", []))}
    if notes:
        result["notes"] = notes
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
