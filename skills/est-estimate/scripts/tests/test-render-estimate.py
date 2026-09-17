#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Tests for render-estimate.py.

The load-bearing property: the interactive HTML must reproduce the engine's arithmetic,
including the project overheads that shrink when scope is cut. A negotiation tool that
shows a saving the client will not actually get is worse than no tool.
"""

import csv
import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from fixtures import feature, inventory, model, options  # noqa: E402

SCRIPTS = Path(__file__).resolve().parent.parent


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / filename)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


est = load("estimate", "estimate.py")
render = load("render_estimate", "render-estimate.py")

try:
    import openpyxl  # noqa: F401
    _HAS_XLSX = True
except ImportError:
    _HAS_XLSX = False


def estimate_for(features=None, **opt):
    inv = inventory(features)
    opt.setdefault("granularity", inv.get("granularity", "project"))
    return est.build_estimate(inv, model(), options(**opt))


class TestMarkdown(unittest.TestCase):
    def test_the_headline_is_the_role_table_and_there_is_no_grand_total(self):
        """Adding up the story rows gave 1,752 h against a 2,414 h headline, because planning,
        QA and overhead touch no story. Summing across roles answers no question anyone asks,
        so the number that could not be reconciled is simply not reported."""
        est = estimate_for()
        md = render.markdown(est)
        self.assertIn("## Hours by role", md)
        self.assertNotIn("hours  ·  range", md)
        for key in ("low", "likely", "high"):
            self.assertNotIn(f"{est['total_hours'][key]:,.0f} hours", md)

    def test_every_role_is_a_range_with_its_two_sources_shown(self):
        md = render.markdown(estimate_for())
        self.assertIn("| Role | Low | Likely | High | Hours | On stories | Project-level | "
                      "Risk | Planned |", md)
        self.assertIn("architect", md)

    def test_every_story_row_carries_a_range_not_a_point(self):
        md = render.markdown(estimate_for([feature("F1", "Login")]))
        self.assertIn("| Low | Likely | High |", md)
        row = [ln for ln in md.splitlines() if ln.startswith("| F1 |")][0]
        cells = [c.strip() for c in row.split("|")]
        lo, likely, hi = (float(cells[7]), float(cells[8]), float(cells[9]))
        self.assertLess(lo, likely)
        self.assertLess(likely, hi)

    def test_every_role_cell_on_a_story_is_a_range_too(self):
        md = render.markdown(estimate_for([feature("F1", "Login")]))
        row = [ln for ln in md.splitlines() if ln.startswith("| F1 |")][0]
        priced = [c.strip() for c in row.split("|") if "–" in c and c.strip()[0].isdigit()]
        self.assertTrue(priced, "no role cell rendered as a range")
        for cell in priced:
            lo, likely, hi = (float(x) for x in cell.split("–"))
            self.assertLessEqual(lo, likely)
            self.assertLessEqual(likely, hi)

    def test_every_feature_row_carries_its_source(self):
        md = render.markdown(estimate_for([feature("F1", "Login")]))
        self.assertIn("S1 §1", md)

    def test_an_uncalibrated_model_is_declared(self):
        """The shipped model is now fitted to one project, so the banner has to be provoked
        rather than assumed — and the thing worth protecting is that a model without a
        calibration history never renders without saying so."""
        e = estimate_for()
        e["cost_model_snapshot"] = {k: v for k, v in e["cost_model_snapshot"].items()
                                    if k != "calibration_history"}
        self.assertIn("Uncalibrated model", render.markdown(e))

    def test_a_calibrated_model_does_not_carry_the_banner(self):
        self.assertNotIn("Uncalibrated model", render.markdown(estimate_for()))

    def test_planning_review_is_called_out_as_the_anchor(self):
        self.assertIn("defended hardest", render.markdown(estimate_for()))

    def test_risk_quadrant_is_surfaced_when_present(self):
        md = render.markdown(estimate_for([
            feature("F1", "Legacy payments bridge", compressibility="low", review_tier="critical")]))
        self.assertIn("Where the risk is", md)
        self.assertIn("Legacy payments bridge", md)

    def test_no_risk_section_when_nothing_is_in_that_quadrant(self):
        # Standing work is off here because environment and release work is genuinely
        # low-compressibility and sensitive, so it always populates the quadrant — which is
        # true, and would make this test about the catalogue rather than about the section.
        md = render.markdown(estimate_for([
            feature("F1", "Marketing page", compressibility="high", review_tier="routine")],
            no_standing_work=True))
        self.assertNotIn("Where the risk is", md)

    def test_the_feature_table_shows_who_does_the_work(self):
        """A row reading "9h" invites a haggle; a row reading "dev 6.1, ba 1.5" invites a
        conversation about who is on it — and the dash against ux is the visible half of
        the fix, since a reader has to be able to see that a role was excluded rather than
        rounded away."""
        md = render.markdown(estimate_for([
            feature("F1", "Nightly reconciliation job", surfaces=["backend"]),
            feature("F2", "Onboarding screens", surfaces=["frontend", "design"])]))
        header = next(l for l in md.split("\n") if l.startswith("| ID | Feature"))
        columns = [c.strip() for c in header.split("|")]
        self.assertIn("dev", columns)
        self.assertIn("ux", columns)
        backend = next(l for l in md.split("\n") if l.startswith("| F1 |"))
        design = next(l for l in md.split("\n") if l.startswith("| F2 |"))
        ux = columns.index("ux")
        self.assertEqual(backend.split("|")[ux].strip(), "—")
        self.assertNotEqual(design.split("|")[ux].strip(), "—")

    def test_standing_work_is_not_labelled_as_something_the_client_can_decline(self):
        md = render.markdown(estimate_for([feature("F1", "Marketing page")]))
        row = next(l for l in md.split("\n") if "SW-ci_pipeline" in l)
        self.assertIn("standing work", row)
        self.assertNotIn("outside agreed scope", row)

    def test_build_compression_is_internal_only(self):
        e = estimate_for()
        self.assertNotIn("Build compression", render.markdown(e, show_manual_baseline=False))
        self.assertIn("Build compression", render.markdown(e, show_manual_baseline=True))

    def test_scope_table_shows_both_apportioned_and_standalone(self):
        md = render.markdown(estimate_for([
            feature("F1", "Agreed"),
            feature("F2", "Extra", scope_status="outside_agreed_scope")]))
        self.assertIn("Share of this project", md)
        self.assertIn("Cost on its own", md)
        self.assertIn("will not save the apportioned figure", md)

    def test_traceability_findings_are_shown_not_hidden(self):
        md = render.markdown(estimate_for([feature("F1", citations=[])]))
        self.assertIn("no citation", md)


class GroupedByEpic(unittest.TestCase):
    """The estimate reads in build order. A flat table of 120 rows answers "what does line 84
    cost"; a reader deciding what to fund first is asking a question only grouping answers."""

    def est(self, **opt):
        inv = inventory([feature("F1", "Screens", epic_id="E2"),
                         feature("F2", "Data model", epic_id="E1")])
        inv["schema_version"] = "1.1"
        # Array order is deliberately the wrong way round: the sequence is what must win.
        inv["epics"] = [{"id": "E2", "name": "Booking", "origin": "source", "sequence": 2,
                         "sequence_why": "reads the data model"},
                        {"id": "E1", "name": "Foundation", "origin": "source", "sequence": 1,
                         "sequence_why": "everything stands on it"}]
        opt.setdefault("granularity", inv.get("granularity", "project"))
        opt.setdefault("no_standing_work", True)
        return est.build_estimate(inv, model(), options(**opt))

    def test_the_features_section_carries_one_heading_per_epic_in_sequence(self):
        md = render.markdown(self.est())
        heads = [l for l in md.split("\n") if l.startswith("### ")]
        self.assertEqual(heads[:2], ["### 1. E1 — Foundation", "### 2. E2 — Booking"])

    def test_each_epic_carries_a_subtotal(self):
        md = render.markdown(self.est())
        self.assertIn("**E1 — Foundation — ", md)
        self.assertIn("h likely**", md)

    def test_the_build_sequence_table_is_present_and_ordered(self):
        md = render.markdown(self.est())
        self.assertIn("## Build sequence", md)
        table = md[md.index("## Build sequence"):md.index("## Features")]
        self.assertLess(table.index("Foundation"), table.index("Booking"))
        self.assertIn("everything stands on it", table)

    def test_the_subtotals_reconcile_with_the_rows_above_them(self):
        """A subtotal that does not add up is worse than none, because it reads as checked."""
        e = self.est()
        md = render.markdown(e)
        section = md[md.index("### 1. E1"):md.index("### 2. E2")]
        rows = [l for l in section.split("\n") if l.startswith("| F")]
        total = sum(float(l.split("|")[8].strip().replace(",", "")) for l in rows)
        stated = float(section.split("h likely**")[0].split("— ")[-1].replace(",", ""))
        self.assertAlmostEqual(total, stated, delta=0.15)

    def test_an_ungrouped_estimate_stays_one_flat_table(self):
        md = render.markdown(estimate_for([feature("F1", "Login")], no_standing_work=True))
        self.assertNotIn("## Build sequence", md)
        self.assertEqual([l for l in md.split("\n") if l.startswith("### ")], [])

    def test_standing_work_lands_in_its_own_section_rather_than_inside_an_epic(self):
        """It is real scope and it is not the client's, so it must be visible and separate."""
        md = render.markdown(self.est(no_standing_work=False))
        self.assertIn("### Ungrouped", md)
        section = md[md.index("### Ungrouped"):]
        self.assertIn("SW-ci_pipeline", section)

    def test_the_brief_carries_the_structure_so_the_agent_can_reason_over_it(self):
        b = render.brief(self.est())
        self.assertEqual([(e["id"], e["sequence"]) for e in b["epics"]], [("E1", 1), ("E2", 2)])
        by_id = {f["id"]: f for f in b["features"]}
        self.assertEqual(by_id["F2"]["epic_sequence"], 1)

    def test_the_spreadsheet_reads_in_the_same_order_as_the_document(self):
        """One field, one answer. Appending standing work put it at the bottom of the sheet
        while the markdown grouped it into the epic the same field named."""
        e = self.est(no_standing_work=False)
        inv = {"features": [], "epics": e["epics"], "sources": [], "implicit_scope": []}
        (_columns, rows), _tasks = render.tabular(e, inv)
        seqs = [r.get("epic_seq") for r in rows if r.get("epic_seq") not in ("", None)]
        self.assertEqual(seqs, sorted(seqs))

    def test_a_standing_row_placed_in_an_epic_sorts_into_it(self):
        e = self.est(no_standing_work=False)
        for f in e["features"]:
            if f.get("origin") == "standing":
                f["epic_id"], f["epic_sequence"] = "E1", 1
        inv = {"features": [], "epics": e["epics"], "sources": [], "implicit_scope": []}
        (_columns, rows), _tasks = render.tabular(e, inv)
        placed = [i for i, r in enumerate(rows) if str(r.get("id") or "").startswith("SW-")]
        others = [i for i, r in enumerate(rows) if r.get("epic_seq") == 2]
        self.assertTrue(max(placed) < min(others), "standing work sorted below a later epic")

    def test_the_brief_carries_what_foundation_work_was_declined(self):
        b = render.brief(self.est(no_standing_work=False))
        self.assertIn("basis", b["standing_work"])


class TestCsv(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def rows(self, e):
        target = self.dir / "estimate.csv"
        render.write_csv(e, target)
        with target.open(encoding="utf-8-sig") as fh:
            return list(csv.DictReader(fh))

    def test_feature_rows_carry_the_cost_model_axes_and_components(self):
        row = self.rows(estimate_for([feature("F1", "Login", review_tier="sensitive")]))[0]
        self.assertEqual(row["review_tier"], "sensitive")
        self.assertTrue(float(row["review"]) > 0)
        self.assertTrue(float(row["build"]) > 0)

    def test_project_components_appear_as_their_own_lines(self):
        names = [r["name"] for r in self.rows(estimate_for())]
        self.assertTrue(any("[project] planning review" in n for n in names))
        self.assertTrue(any("[project] qa" in n for n in names))

    def test_no_grand_total_row_is_written(self):
        names = [r["name"] for r in self.rows(estimate_for())]
        for gone in ("[total] likely", "[total] low", "[total] high"):
            self.assertNotIn(gone, names)

    def test_every_row_carries_its_range_and_every_role_carries_one(self):
        est = estimate_for([feature("F1", "Login")])
        rows = {r["id"]: r for r in self.rows(est)}
        row = rows["F1"]
        self.assertLess(float(row["hours_low"]), float(row["hours_high"]))
        roles = [r for r in est["by_role"] if f"{r}_low" in row]
        self.assertTrue(roles, "no per-role columns in the sheet")
        for role in roles:
            self.assertLessEqual(float(row[f"{role}_low"]), float(row[f"{role}_high"]))


class ExtendsTheInventory(unittest.TestCase):
    """The estimate sheets are the inventory sheets with more columns on them.

    The old estimate.csv re-derived its own 23 columns and lost description, epic name,
    surfaces, task ids, open questions, locations and quotes — and every task. The columns
    here come from est-scope-extract's own row builders, so a column added there arrives
    without anyone remembering to add it.
    """

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.inv = inventory([feature("F1", "Login")])
        self.inv["epics"] = [{"id": "E1", "name": "Access", "origin": "source"}]
        self.inv["features"][0]["epic_id"] = "E1"
        self.inv["features"][0]["tasks"] = [
            {"id": "F1-T1", "name": "Sign in",
             "citations": [{"source_id": "S1", "location": "§1", "quote": "Users must log in."}]}]
        self.est = est.build_estimate(self.inv, model(), options(granularity="project"))

    def tearDown(self):
        self.tmp.cleanup()

    def inventory_columns(self):
        inv_mod = render.inventory_renderer()
        return list(inv_mod.STORY_COLUMNS), list(inv_mod.TASK_COLUMNS)

    def test_no_inventory_story_column_is_dropped(self):
        story_columns, _ = self.inventory_columns()
        target = self.dir / "estimate.csv"
        render.write_csv(self.est, target, self.inv)
        with target.open(encoding="utf-8-sig") as fh:
            got = next(csv.reader(fh))
        for column in story_columns:
            self.assertIn(column, got, f"est-scope-extract wrote {column} and the estimate lost it")

    def test_no_inventory_task_column_is_dropped(self):
        _, task_columns = self.inventory_columns()
        target = self.dir / "estimate.tasks.csv"
        render.write_tasks_csv(self.est, target, self.inv)
        with target.open(encoding="utf-8-sig") as fh:
            got = next(csv.reader(fh))
        for column in task_columns:
            self.assertIn(column, got)

    def test_the_inventory_content_arrives_not_just_the_headers(self):
        target = self.dir / "estimate.csv"
        render.write_csv(self.est, target, self.inv)
        with target.open(encoding="utf-8-sig") as fh:
            row = next(iter(csv.DictReader(fh)))
        self.assertEqual(row["epic_name"], "Access")
        self.assertEqual(row["description"],
                         self.inv["features"][0]["description"],
                         "the description est-scope-extract wrote did not survive the join")
        self.assertEqual(row["task_ids"], "F1-T1")

    def test_a_task_id_a_story_names_resolves_to_a_row(self):
        render.write_csv(self.est, self.dir / "estimate.csv", self.inv)
        render.write_tasks_csv(self.est, self.dir / "estimate.tasks.csv", self.inv)
        with (self.dir / "estimate.csv").open(encoding="utf-8-sig") as fh:
            story = next(iter(csv.DictReader(fh)))
        with (self.dir / "estimate.tasks.csv").open(encoding="utf-8-sig") as fh:
            tasks = {r["task_id"] for r in csv.DictReader(fh)}
        for tid in story["task_ids"].split("; "):
            self.assertIn(tid, tasks)

    def test_a_story_priced_but_absent_from_the_inventory_still_appears(self):
        """Standing work is priced and is in no inventory. Dropping it from the sheet would
        price less on the page than the estimate does."""
        target = self.dir / "estimate.csv"
        render.write_csv(self.est, target, self.inv)
        with target.open(encoding="utf-8-sig") as fh:
            ids = {r["id"] for r in csv.DictReader(fh)}
        priced = {f["id"] for f in self.est["features"]}
        self.assertTrue(priced <= ids, sorted(priced - ids))

    def test_without_the_inventory_it_still_renders_and_says_what_is_missing(self):
        est_copy = json.loads(json.dumps(self.est))
        est_copy["inventory"] = "nowhere/feature-inventory.json"
        inv, problem = render.load_inventory(est_copy)
        self.assertIsNone(inv)
        self.assertIn("inventory not found", problem)
        target = self.dir / "estimate.csv"
        render.write_csv(est_copy, target, None)
        with target.open(encoding="utf-8-sig") as fh:
            row = next(iter(csv.DictReader(fh)))
        self.assertEqual(row["id"], "F1")
        self.assertTrue(float(row["hours_high"]) > 0)


@unittest.skipUnless(_HAS_XLSX, "openpyxl not importable")
class Workbook(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.inv = inventory([feature("F1", "Login")])
        self.inv["features"][0]["tasks"] = [
            {"id": "F1-T1", "name": "Sign in",
             "citations": [{"source_id": "S1", "location": "§1", "quote": "Users must log in."}]}]
        self.est = est.build_estimate(self.inv, model(), options(granularity="project"))
        self.target = self.dir / "estimate.xlsx"
        self.assertTrue(render.write_xlsx(self.est, self.target, self.inv))

    def tearDown(self):
        self.tmp.cleanup()

    def test_the_same_two_tabs_the_inventory_workbook_has(self):
        from openpyxl import load_workbook
        self.assertEqual(load_workbook(self.target).sheetnames, ["Stories", "Tasks"])

    def test_the_priced_columns_are_on_the_stories_tab_as_ranges(self):
        from openpyxl import load_workbook
        ws = load_workbook(self.target)["Stories"]
        header = [c.value for c in ws[1]]
        for column in ("hours_low", "hours_likely", "hours_high"):
            self.assertIn(column, header)
        for role in self.est["by_role"]:
            if f"{role}_low" in header:
                break
        else:
            self.fail("no per-role columns on the workbook")

    def test_a_story_still_clicks_through_to_its_tasks(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.target)
        ws, ts = wb["Stories"], wb["Tasks"]
        col = [c.value for c in ws[1]].index("task_ids") + 1
        cell = ws.cell(row=2, column=col)
        self.assertIsNotNone(cell.hyperlink)
        at = int(cell.hyperlink.location.split("!A")[1])
        self.assertEqual(ts.cell(row=at, column=1).value, "F1-T1")

    def test_every_link_carries_the_display_text_google_sheets_needs(self):
        """est-estimate writes these tabs through est-scope-extract's own `write_workbook`, so
        this asserts the same claim on this side of the boundary: without `display`, Google
        Sheets shows the address instead of the name. Excel never did, which is why a whole
        column of `#gid=1123955261&range=A2` reached a reader before anyone noticed."""
        from openpyxl import load_workbook
        wb = load_workbook(self.target)
        links = [(n, c) for n in wb.sheetnames for row in wb[n].iter_rows()
                 for c in row if c.hyperlink]
        self.assertTrue(links, "the workbook has no links to check")
        for name, cell in links:
            self.assertEqual(getattr(cell.hyperlink, "display", None), str(cell.value),
                             f"{name}!{cell.coordinate} would show its address in Sheets")

    def test_without_openpyxl_the_render_still_succeeds_and_says_the_workbook_was_skipped(self):
        saved = sys.modules.get("openpyxl")
        sys.modules["openpyxl"] = None
        try:
            self.assertFalse(render.write_xlsx(self.est, self.dir / "x.xlsx", self.inv))
        finally:
            if saved is None:
                sys.modules.pop("openpyxl", None)
            else:
                sys.modules["openpyxl"] = saved


class TestBrief(unittest.TestCase):
    """What a conversational agent needs to defend a number, without the ledger payload."""

    def test_every_tag_keeps_the_reason_it_was_given(self):
        f = feature("F1", "Refunds", review_tier="sensitive")
        f["tags"]["review_tier"]["why"] = "handles refunds against a customer card"
        b = render.brief(estimate_for([f]))
        self.assertEqual(b["features"][0]["why"]["review_tier"]["why"],
                         "handles refunds against a customer card")

    def test_the_verbatim_quote_survives_to_the_brief(self):
        b = render.brief(estimate_for([feature("F1", "Login")]))
        self.assertIn("Login is required", b["features"][0]["source"]["quote"])

    def test_the_dominant_cost_driver_is_named_per_feature(self):
        b = render.brief(estimate_for([feature("F1", "Payments", review_tier="critical")]))
        self.assertEqual(b["features"][0]["dominant_component"], "review")

    def test_the_brief_can_answer_who_does_this_work(self):
        """Nadia's rule is that no number is asserted that cannot be decomposed on demand.
        A per-story figure with no role split decomposes to hours and stops exactly where
        the question usually goes next."""
        b = render.brief(estimate_for([feature("F1", "Reconciliation", surfaces=["backend"])]))
        row = next(f for f in b["features"] if f["id"] == "F1")
        self.assertEqual(sorted(row["by_role"]), ["ba", "dev"])

    def test_the_brief_drops_the_cost_model_snapshot(self):
        b = render.brief(estimate_for())
        self.assertNotIn("cost_model_snapshot", b)
        self.assertIn("calibrated", b)

    def test_the_brief_is_far_smaller_than_the_ledger_payload(self):
        e = estimate_for([feature(f"F{i}") for i in range(1, 6)])
        self.assertLess(len(json.dumps(render.brief(e))), len(json.dumps(e)) / 3)


class TestInteractiveHtml(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def build(self, features=None):
        e = estimate_for(features)
        target = self.dir / "estimate.html"
        render.write_html(e, target, {"stack": "standard_saas", "qa_platform": "web",
                                      "engagement": "standard"})
        return e, target.read_text(encoding="utf-8")

    def embedded(self, html):
        start = html.index('<script id="estimate-data" type="application/json">') + \
            len('<script id="estimate-data" type="application/json">')
        end = html.index("</script>", start)
        return json.loads(html[start:end].replace("<\\/", "</"))

    def test_the_estimate_survives_embedding_intact(self):
        e, html = self.build()
        self.assertAlmostEqual(self.embedded(html)["total_hours"]["likely"],
                               e["total_hours"]["likely"])

    def test_the_cost_model_travels_with_the_page_so_it_can_recompute(self):
        _, html = self.build()
        data = self.embedded(html)
        self.assertIn("planning", data["cost_model_snapshot"])
        self.assertIn("overhead_rate", data["cost_model_snapshot"])
        self.assertIn("_render_options", data)

    def test_a_closing_script_tag_in_the_data_cannot_break_the_page(self):
        """A feature named with a closing tag would otherwise end the script element early."""
        _, html = self.build([feature("F1", "Report </script> injection")])
        self.assertNotIn("</script> injection", html)
        self.assertEqual(self.embedded(html)["features"][0]["name"], "Report </script> injection")

    def test_the_page_carries_its_own_recompute_self_check(self):
        _, html = self.build()
        self.assertIn("Recompute mismatch", html)

    def test_the_self_check_covers_the_whole_interval_not_just_the_mean(self):
        """A mean-only check is blind to band divergence, which is the module's headline claim
        and the thing that actually drifted."""
        _, html = self.build()
        self.assertIn("['low', 'likely', 'high']", html)

    def test_citations_reach_the_page_for_every_feature(self):
        _, html = self.build([feature("F1", "Login")])
        self.assertEqual(self.embedded(html)["features"][0]["citations"][0]["location"], "§1")




class TheBudgetColumnsAddUp(unittest.TestCase):
    """The per-role risk and planned block, restored on 2026-09-16.

    It was added, then removed by a commit whose subject was about something else, and the
    removal missed three files — `check-parity.py` and its test kept comparing fields nothing
    emitted any more, so the parity check passed by skipping every comparison, and the module's
    own outputs contract went on promising columns that did not exist. These tests exist so the
    next removal has to be deliberate.

    The claim that makes the columns worth having is that they ADD UP: a band does not, because
    variances combine in quadrature, but systematic risk is correlated by definition and splits
    linearly. A reader can sum the Risk column down the sheet and get the project's figure.
    """

    def setUp(self):
        self.est = estimate_for()

    def test_every_role_row_reconciles_with_itself(self):
        for role, row in self.est["by_role"].items():
            with self.subTest(role=role):
                self.assertAlmostEqual(row["hours"],
                                       row["on_stories"] + row["project_level"], delta=0.15)
                self.assertAlmostEqual(row["risk_adjusted_hours"],
                                       row["hours"] + row["model_risk_hours"], delta=0.15)

    def test_planned_is_narrower_than_the_band_it_sits_beside(self):
        """It is a budget line, not a worst case. Quoting it as the top of the range is the
        specific misreading the column's own note exists to prevent."""
        for role, row in self.est["by_role"].items():
            with self.subTest(role=role):
                self.assertLess(row["risk_adjusted_hours"], row["high"] + 0.15)

    def test_the_risk_column_sums_down_the_sheet(self):
        """The property a budget needs and a band cannot give."""
        total = sum(r["model_risk_hours"] for r in self.est["by_role"].values())
        attribution = self.est["confidence"].get("role_attribution") or {}
        self.assertTrue(attribution, "role_attribution is how the reconciliation is published")
        self.assertAlmostEqual(total, attribution.get("model_risk_hours", total), delta=0.3)

    def test_the_csv_carries_the_budget_block_for_roles_that_work_on_stories(self):
        """Only those roles. QA is priced as a share of the story total and the architect on the
        calendar, so neither sits on a story — their columns were zero on every story row with
        the real figure one row below on `[project] qa`. Twelve columns of zeros read as hours
        gone missing, and were read that way."""
        columns = render.estimate_columns(self.est)
        working = render.story_roles(self.est)
        self.assertTrue(working, "no role carries story hours in the fixture")
        for role in self.est["by_role"]:
            for suffix in ("_low", "_likely", "_high", "_hours", "_risk", "_planned"):
                if role in working:
                    self.assertIn(f"{role}{suffix}", columns)
                else:
                    self.assertNotIn(f"{role}{suffix}", columns)
        self.assertIn("risk_hours", columns)
        self.assertIn("planned_hours", columns)

    def test_a_project_level_role_gets_no_story_columns_but_keeps_its_hours(self):
        """Removing a column must not remove a number. The `[project]` row carries the same
        figures its role columns used to duplicate, and says so in its own description."""
        absent = [r for r in self.est["by_role"] if r not in render.story_roles(self.est)]
        self.assertTrue(absent, "the fixture prices no role at project level")
        (columns, rows), _ = render.tabular(self.est)
        for role in absent:
            self.assertNotIn(f"{role}_hours", columns)
            line = next(r for r in rows if r.get("name") == f"[project] {role}")
            self.assertAlmostEqual(line["hours"],
                                   self.est["by_role"][role]["hours"], delta=0.15)
            self.assertEqual(line["risk_hours"],
                             self.est["project_components"][role]["model_risk_hours"])
            self.assertIn("Priced per project, not per story", line["description"])
            self.assertIn(f"{role}_hours", line["description"])

    def test_the_rule_is_the_data_not_a_list_of_role_names(self):
        """Keyed on `on_stories`, so a cost model that puts QA on stories gets the columns back
        with no code change — and a project with no infra stories loses `devops_*`, which is the
        same noise for the same reason."""
        moved = json.loads(json.dumps(self.est))
        moved["by_role"]["qa"]["on_stories"] = 12.0
        self.assertIn("qa", render.story_roles(moved))
        self.assertIn("qa_hours", render.estimate_columns(moved))
        moved["by_role"]["devops"]["on_stories"] = 0.0
        self.assertNotIn("devops", render.story_roles(moved))
        self.assertNotIn("devops_hours", render.estimate_columns(moved))

    def test_the_task_sheet_drops_the_same_roles(self):
        """`alloc_qa_*` was zero on every task for the same reason: allocation splits a STORY's
        hours, and QA sits on no story."""
        columns = render.task_price_columns(self.est)
        for role in self.est["by_role"]:
            present = f"alloc_{role}_hours" in columns
            self.assertEqual(present, role in render.story_roles(self.est), role)

    def test_the_estimate_declares_how_task_figures_were_allocated(self):
        """Tasks are source rows, not estimable units, and no anchor carries per-task effort.
        The split is therefore the rule that asserts least — and it is named rather than
        implied, so nobody negotiates over a task figure as though it were estimated."""
        allocation = self.est.get("task_allocation")
        self.assertTrue(allocation)
        self.assertIn("even", json.dumps(allocation).lower())


class TheEstimateDoesNotAnswerHowLongItTakes(unittest.TestCase):
    """That job moved to est-plan, and the move is the fix rather than the wording.

    The report used to print a span derived from `hours / six people` under a heading that
    called it derived and not a commitment, and clients quoted it back as a schedule. The span
    is still computed — the architect and ceremony lines are priced against it — and its basis
    is still stated as the premise of those numbers. What is gone is the figure.
    """

    def setUp(self):
        self.est = estimate_for()

    def test_the_markdown_offers_no_week_count_as_an_answer(self):
        text = render.markdown(self.est)
        self.assertNotIn("## Calendar duration", text)
        self.assertIn("## How long will it take?", text)
        self.assertIn("/est-plan", text)
        weeks = self.est.get("duration", {}).get("weeks")
        if weeks:
            self.assertNotIn(f"**{weeks} weeks.**", text)

    def test_the_basis_survives_because_it_is_a_premise_and_not_an_answer(self):
        """Removing it would hide what the architect and overhead lines rest on, which the
        module's traceability bar does not allow."""
        joined = " ".join(self.est["assumptions"])
        self.assertIn("ceremony", joined)
        self.assertIn("NOMINAL", joined)
        self.assertIn("/est-plan", joined)

    def test_the_html_panel_points_at_the_planner_instead_of_printing_a_figure(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "estimate.html"
            render.write_html(self.est, target, {})
            html = target.read_text(encoding="utf-8")
        self.assertIn("How long will it take?", html)
        self.assertNotIn("Calendar duration — derived, not a commitment", html)
        self.assertIn("/est-plan", html)

    def test_the_report_carries_the_brand_rather_than_a_palette_of_its_own(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "estimate.html"
            render.write_html(self.est, target, {})
            html = target.read_text(encoding="utf-8")
        self.assertNotIn("__BRAND_CSS__", html)
        self.assertIn("--brand: #002c8d", html)


if __name__ == "__main__":
    unittest.main()
