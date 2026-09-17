#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Project plan.json into the documents a person reads: markdown, and Gantt tabs.

The workbook is EXTENDED, never replaced. `est-estimate` already wrote `estimate.xlsx` with
its Stories and Tasks tabs and the hyperlinks between them; this opens that file and appends
sheets to it, so a client has one workbook rather than two that have to be read side by side.
It deliberately does not touch `render-inventory.py:write_workbook`, which est-scope-extract
and est-estimate share — putting a Gantt concern inside a function two other skills call would
make a scheduling bug their problem too.

The Gantt is a CELL GRID, not a chart object. Readability is the requirement, and a grid
prints, survives a copy-paste into a deck, collapses by outline level, and can carry a
hyperlink on every row — none of which an embedded chart does. One column per week, labelled
W1..Wn, because this module's plans carry no dates.

    uv run scripts/render-plan.py <plan.json> --workbook estimate.xlsx [-o plan.md]
"""

import argparse
import importlib.util
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
BRAND = HERE.parent.parent / "est-estimate" / "scripts" / "brand.py"

ROLE_ORDER = ("ba", "ux", "architect", "dev", "qa", "devops")
ROLE_LABEL = {"ba": "BA", "ux": "UX", "dev": "Dev", "qa": "QA",
              "devops": "DevOps", "architect": "Architect"}


def brand_module():
    spec = importlib.util.spec_from_file_location("brand", BRAND)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def link_label(value):
    """The text a hyperlink shows, for the `display` attribute — see est-scope-extract's
    render-inventory.link_label for why it is needed and why the 255-character cap is there.

    Deliberately duplicated rather than imported: the two skills install independently, and a
    cross-skill import for three lines would make this one unable to run without the other.
    """
    text = "" if value is None else str(value)
    return text[:252] + "..." if len(text) > 255 else text


# --- markdown ----------------------------------------------------------------

def markdown(plan):
    """Every option, its hours, its score and its risks — the recommendation last, explained.

    Options are presented complete. A schedule without its cost is half an answer, and the
    module's rule is that the reader compares plan-and-cost pairs rather than calendars.
    """
    out = [f"# Delivery options — {plan.get('project') or 'untitled'}", "",
           "Each option below is a complete plan **and its own estimate**. The same scope "
           "sequenced and staffed differently costs differently: the architect is priced as "
           "setup plus a weekly presence, and ceremony is priced per person per week, so a "
           "longer plan genuinely costs more and a leaner one genuinely costs less. Story "
           "hours never move.", "",
           "**Weeks are relative.** Week 1 is whenever this starts. This plan carries no dates.",
           ""]

    roster = plan["inputs"].get("roster") or {}
    out += ["## Team", "",
            ("Headcount is an output here, not a constraint — the sweep explores team shapes "
             "and every addition below is justified or refused against the backlog's own "
             "dependency graph."), ""]
    if roster:
        out += [f"You supplied: {', '.join(f'{n}x {ROLE_LABEL.get(r, r)}' for r, n in sorted(roster.items()))}. "
                "Treated as a starting point, not a ceiling.", ""]
    for row in plan["staffing"].get("oversized_roster") or []:
        out += [f"> **The roster is larger than this backlog can occupy.** {row['note']} "
                f"On this scope {row['sustainable_count']} would be busy.", ""]

    allowed = [d for d in plan["staffing"]["decisions"] if d["allowed"]]
    refused = [d for d in plan["staffing"]["decisions"] if not d["allowed"]]
    if allowed:
        out += ["**Additions this backlog justifies**", ""]
        for d in allowed:
            for line in d["justification"]:
                out.append(f"- {line}" if not line.startswith("  ") else f"  - {line.strip()}")
        out.append("")
    if refused:
        out += ["**Additions it refuses**", ""]
        for d in refused:
            out.append(f"- **{ROLE_LABEL.get(d['role'], d['role'])} → {d['count']}.** "
                       + " ".join(d["refusals"]))
        out.append("")

    order = plan.get("build_order") or {}
    if order.get("epics"):
        out += ["## Build order", "",
                "What each epic stands on, and why. Every schedule below respects this — the "
                "same graph gates the calendar, so a bar cannot start before what it depends "
                "on, whatever the archetype.", "",
                order.get("how", ""), "",
                "| # | Epic | Waits for | Why |", "| ---: | --- | --- | --- |"]
        for row in order["epics"]:
            waits = row.get("waits_for") or []
            out.append(f"| {row.get('sequence')} | {row.get('name') or row['epic_id']} | "
                       + (", ".join(w["epic_id"] for w in waits) or "—") + " | "
                       + (waits[-1]["why"] if waits else "nothing precedes it") + " |")
        out.append("")
        for row in order.get("cycles_broken") or []:
            out += [f"> **{row['epic']} and {row['waits_on']} depend on each other.** "
                    f"{row['dropped']} The dropped edge was: {row['why']}.", ""]

    audit = plan.get("ordering_check") or {}
    if audit and not audit.get("ok"):
        out += ["> **This plan's calendar contradicts its own dependency graph.** "
                + f"{len(audit.get('findings') or [])} orderings are violated. "
                  "It is printed so it can be argued with, and it is not shippable.", ""]
        for finding in (audit.get("findings") or [])[:8]:
            out.append(f"> - {finding['detail']}")
        out.append("")

    multi = any(len(o.get("phases") or []) > 1 for o in plan["options"])
    out += ["## The options", ""]
    if multi:
        out += ["Phase 1 first, because a later phase is often a separate agreement and the "
                "number a client signs should not silently contain it. The whole-programme "
                "columns are what the option actually delivers, and the fit score is computed "
                "on them — every option holds the same scope, so scoring phase 1 alone would "
                "let an option win by pushing work across the boundary.", "",
                "| # | Option | P1 weeks | P1 h | Whole weeks | Whole h | Range | Fit |",
                "| --- | --- | ---: | ---: | ---: | ---: | --- | ---: |"]
    else:
        out += ["| # | Option | Weeks | Likely h | Range | Fit |",
                "| --- | --- | ---: | ---: | --- | ---: |"]
    for i, o in enumerate(plan["options"], start=1):
        t = o["estimate"]["total_hours"]
        mark = " ★" if o["id"] == plan.get("recommended") else ""
        flag = "" if o["feasibility"]["feasible"] else " ⚠ not deliverable"
        head = (f"| {i}{mark} | {o['archetype'].title()}, "
                f"{team_label(o['team_shape'])}{flag} | ")
        phases = o.get("phases") or []
        if multi:
            first = phases[0] if phases else None
            out.append(head
                       + (f"{first['finishes_week']:.1f} | {first['standalone_hours']:.0f} | "
                          if first else "— | — | ")
                       + f"{o['schedule']['weeks']:.1f} | {t['likely']:.0f} | "
                         f"{t['low']:.0f}–{t['high']:.0f} | {o['score']['fit']:.1f} |")
        else:
            out.append(head + f"{o['schedule']['weeks']:.1f} | {t['likely']:.0f} | "
                              f"{t['low']:.0f}–{t['high']:.0f} | {o['score']['fit']:.1f} |")
    out.append("")

    baseline = plan.get("baseline")
    if baseline is not None and not any(o["id"] == baseline["id"] for o in plan["options"]):
        # Only when the sweep never considered it, which means a supplied roster put its floor
        # above one person per role. It is printed as a reference, not offered as an option —
        # a roster is a floor, and proposing a smaller team than the one the user says they
        # have would be answering a question nobody asked.
        out += ["## The baseline", "",
                baseline["baseline_note"], ""]
        out += option_section(baseline, False)

    for o in plan["options"]:
        out += option_section(o, o["id"] == plan.get("recommended"))

    if plan.get("recommended"):
        best = next(o for o in plan["options"] if o["id"] == plan["recommended"])
        out += ["## Why this one", "",
                f"**{best['archetype'].title()}, {team_label(best['team_shape'])}** scores "
                f"{best['score']['fit']:.1f}. Every sub-score is printed above it; the weights "
                f"live in the cost model's `staffing.scoring` block, so anyone who weighs these "
                f"differently can re-run it and get their own answer rather than arguing with "
                f"this one.", ""]
    else:
        out += ["## No option is deliverable", "",
                "Every schedule considered failed its feasibility check. The failures are "
                "printed above; they are not presentational, and none of these plans should be "
                "shown to a client.", ""]
    return "\n".join(out) + "\n"


def team_label(shape):
    return ", ".join(f"{n}x {ROLE_LABEL.get(r, r)}" for r, n in sorted(shape.items()))


def phase_lead(o):
    """The headline, led by the FIRST phase rather than by the whole programme.

    A later phase is frequently a separate agreement, so a client being quoted an MVP should
    not be handed a span and a total that silently include work under a different contract.
    The whole-programme figure stays on the line beside it, labelled, because it is what the
    option actually contains.
    """
    phases = o.get("phases") or []
    t = o["estimate"]["total_hours"]
    whole = (f"{o['schedule']['weeks']:.1f} weeks, {t['likely']:.0f} h likely "
             f"({t['low']:.0f}–{t['high']:.0f})")
    fit = f"Fit {o['score']['fit']:.1f}/10."
    if len(phases) < 2:
        return f"**{whole}.** {fit}"
    first = phases[0]
    later = len(phases) - 1
    return (f"**Phase {first['phase']}: {first['finishes_week']:.1f} weeks, "
            f"{first['standalone_hours']:.0f} h on its own.** "
            f"Whole programme, including {later} later phase{'s' if later > 1 else ''}: "
            f"{whole}. {fit}")


def option_section(o, recommended):
    t = o["estimate"]["total_hours"]
    out = [f"### {o['archetype'].title()} — {team_label(o['team_shape'])}"
           + ("  ★ recommended" if recommended else ""), "",
           o["what_it_is"], "",
           phase_lead(o), "",
           "| Role | Low | Likely | High |", "| --- | ---: | ---: | ---: |"]
    for role in ROLE_ORDER:
        row = o["estimate"]["by_role"].get(role)
        if row:
            out.append(f"| {ROLE_LABEL.get(role, role)} | {row['low']:.0f} | "
                       f"{row['likely']:.0f} | {row['high']:.0f} |")
    phases = o.get("phases") or []
    if len(phases) > 1:
        out += ["", "**The phases.** A phase is a delivery commitment, not a dependency: "
                "nothing in a later phase is built until the earlier one closes, whether or not "
                "the graph requires it. Analysis may overlap — the next phase can be specified "
                "while this one builds — which is why a phase's window starts before the one "
                "before it ends.", "",
                "| Phase | Runs | On its own | Share of this programme | Why it is a phase |",
                "| ---: | --- | ---: | ---: | --- |"]
        for row in phases:
            out.append(f"| {row['phase']} | W{int(row['starts_week']) + 1}–"
                       f"W{max(1, -(-row['finishes_week'] // 1)):.0f} | "
                       f"{row['standalone_hours']:.0f} h | "
                       f"{row['apportioned_hours']:.0f} h | {row['basis']} |")
        out += ["", "*On its own* is what that phase costs if it is the only thing commissioned; "
                "*share of this programme* is its portion of the total above. The second is "
                "smaller because environments and most planning are paid once however many "
                "phases run — quote the first to anyone deciding whether to commission a phase.",
                ""]

    staggered = [p for p in o["schedule"]["team"] if p.get("join_week", 0.0) > 0.005]
    if staggered:
        out += ["", "**Who arrives when.** Nobody joins before their role has more ready work "
                "than the people already on it can clear — so the plan does not start four "
                "developers on the same Monday against a codebase that does not exist yet.", "",
                "| Person | Joins | First delivery | Hours | Occupied |",
                "| --- | ---: | ---: | ---: | ---: |"]
        for person in o["schedule"]["team"]:
            out.append(f"| {person['name']} | W{int(person.get('join_week', 0.0)) + 1} | "
                       f"W{int(person['starts_week']) + 1} | {person['delivered_hours']:.0f} | "
                       f"{person['utilisation']:.0%} |")
    out += ["", "| Sub-score | /10 | What it rests on |", "| --- | ---: | --- |"]
    for name, part in o["score"]["parts"].items():
        out.append(f"| {name.title()} | {part['score']:.1f} | {part['rests_on']} |")
    out += ["", "**Risks**", ""]
    for row in o["risks"]:
        out.append(f"- *{row['risk']}* — {row['mitigation']}")
    feas = o["feasibility"]
    out += ["", "**Deliverable?** " + ("Yes. " + "; ".join(feas["checks"]) + "."
                                       if feas["feasible"] else
                                       "**No.** " + " ".join(feas["failures"])), ""]
    return out


# --- the Gantt ---------------------------------------------------------------

# Work that belongs to no epic but is real time on somebody's calendar. Naming these
# "unassigned" tells a reader nothing and reads as a mistake — each one has an actual name.
PSEUDO_EPIC = {"planning": "Planning artefacts", "setup": "Project setup",
               "regression": "Final regression pass", "—": "Other"}
COMPONENT_LABEL = {"spec": "Specification", "build": "Build", "review": "Review",
                   "rework": "Rework", "qa": "QA pass", "planning": "Planning"}


def story_rows(option, epics):
    """The chart's tree: epic → story → component slice, from the option's own bookings.

    Rows are STORIES. The first version drew one row per person per epic and rolled the stories
    inside it into a count — "Foundation & DevOps · 5 stories" — which meant the finest thing a
    reader could see was an epic. A schedule is a statement about work items, and somebody
    asking when the Stripe Connect onboarding is built, and what it is waiting for, could not
    answer it from that chart at all.

    (The epic-packet version existed because an even earlier one cut a row every time a person's
    run of work broke, producing 381 rows with one epic appearing four times under one name. The
    right answer to that was to group by epic and draw the stories; collapsing to epic rows threw
    away the thing being scheduled.)

    A story appears ONCE, whatever number of people touch it, with the window it actually
    occupies on the calendar and the people who carry it. Planning and standing work belong to
    no epic and get pseudo-epics of their own — calling them "unassigned" tells a reader nothing
    and reads as a mistake.
    """
    name_of = {e.get("id"): e.get("name") for e in epics}
    sequence = {e.get("id"): e.get("sequence") or 0 for e in epics}
    tree = {}
    for person in option["schedule"]["team"]:
        for item in person["items"]:
            story = item.get("story_id")
            key = item["epic_id"] or ("planning" if item["component"] == "planning"
                                      else "setup" if str(story or "").startswith("SW-")
                                      # The regression pass sweeps the whole scope after the
                                      # last build, so it belongs to every epic and therefore
                                      # to none. It is the only QA booking without one.
                                      else "regression" if item["component"] == "qa"
                                      else "—")
            epic = tree.setdefault(key, {"epic_id": key, "name": name_of.get(key)
                                         or PSEUDO_EPIC.get(key, key),
                                         "sequence": sequence.get(key, 10 ** 6),
                                         "hours": 0.0, "start_week": item["start_week"],
                                         "finish_week": item["finish_week"], "stories": {}})
            epic["hours"] += item["hours"]
            epic["start_week"] = min(epic["start_week"], item["start_week"])
            epic["finish_week"] = max(epic["finish_week"], item["finish_week"])

            # Epic-level work — the QA passes — carries no story id. It is real booked time and
            # it gets a row of its own under the epic rather than vanishing into the band.
            sid = story or f"{key}:{item['component']}"
            row = epic["stories"].setdefault(sid, {
                "story_id": story, "key": sid, "label": item["label"],
                "hours": 0.0, "start_week": item["start_week"],
                "finish_week": item["finish_week"], "by_role": {}, "slices": {}})
            row["hours"] += item["hours"]
            row["start_week"] = min(row["start_week"], item["start_week"])
            row["finish_week"] = max(row["finish_week"], item["finish_week"])
            row["by_role"][person["role"]] = (row["by_role"].get(person["role"], 0.0)
                                              + item["hours"])
            slice_ = row["slices"].setdefault(item["component"], {
                "component": item["component"], "hours": 0.0, "owners": set(),
                "start_week": item["start_week"], "finish_week": item["finish_week"]})
            slice_["hours"] += item["hours"]
            slice_["owners"].add(person["name"])
            slice_["start_week"] = min(slice_["start_week"], item["start_week"])
            slice_["finish_week"] = max(slice_["finish_week"], item["finish_week"])

    for epic in tree.values():
        for row in epic["stories"].values():
            # The bar takes the colour of the role carrying most of the story, so a chart read
            # at a distance still says what KIND of work each row mostly is.
            row["role"] = max(row["by_role"].items(), key=lambda kv: kv[1])[0] \
                if row["by_role"] else None
            row["owners"] = sorted({o for s in row["slices"].values() for o in s["owners"]})
            row["slices"] = sorted(row["slices"].values(),
                                   key=lambda s: (s["start_week"], s["component"]))
        epic["stories"] = sorted(epic["stories"].values(),
                                 key=lambda r: (r["start_week"], str(r["key"])))
    return sorted(tree.values(), key=lambda e: (e["sequence"], e["start_week"]))


def span_weeks(start, finish, weeks):
    """The week columns a piece of work touches. Work shorter than a week still owns the week
    it happened in — an empty row reads as nothing scheduled, which would be a lie."""
    lo = max(0, int(start))
    hi = max(lo, min(weeks - 1, int(finish - 1e-9)))
    return set(range(lo, min(hi, weeks - 1) + 1))


def story_index(workbook):
    """id -> row on the Stories tab, so every Gantt row can link to the work it draws."""
    if "Stories" not in workbook.sheetnames:
        return {}
    sheet = workbook["Stories"]
    header = [c.value for c in sheet[1]]
    if "id" not in header:
        return {}
    column = header.index("id") + 1
    return {sheet.cell(row=r, column=column).value: r
            for r in range(2, sheet.max_row + 1)
            if sheet.cell(row=r, column=column).value}


def gates(plan, option, estimate_features=()):
    """What each epic and each STORY stands on, and the week the last of it clears.

    So a reader can answer "why is this here" off the chart itself. The epic graph comes from
    `plan["build_order"]`, which plan.py wrote from the same `epic_predecessors` the scheduler
    was given — quoting a different graph on the chart from the one the calendar was built
    against is the failure that work exists to remove.

    Rows are stories now, so a story answers with its OWN `depends_on` where it has one, which
    is both more specific and more useful than its epic's predecessors: it names the actual
    work in the way. It falls back to the epic's gate when the story records none of its own.

    Returns `{key: text}` where key is an epic id or a story id.
    """
    epic_done = {"spec": {}, "build": {}}
    story_done = {}
    for person in option["schedule"]["team"]:
        for item in person["items"]:
            key, component, story = item.get("epic_id"), item.get("component"), item.get("story_id")
            if key and component in epic_done:
                epic_done[component][key] = max(epic_done[component].get(key, 0.0),
                                                item["finish_week"])
            if story and component == "build":
                story_done[story] = max(story_done.get(story, 0.0), item["finish_week"])

    out = {}
    for row in (plan.get("build_order") or {}).get("epics") or []:
        waits = [w["epic_id"] for w in row.get("waits_for") or []]
        if not waits:
            continue
        clears = max((epic_done["build"][w] for w in waits if w in epic_done["build"]),
                     default=None)
        out[row["epic_id"]] = (", ".join(waits)
                               + (f" (clears W{max(1, -(-clears // 1)):.0f})" if clears else ""))

    for feature in estimate_features or ():
        deps = [d for d in (feature.get("depends_on") or []) if d in story_done]
        if not deps:
            continue
        clears = max(story_done[d] for d in deps)
        shown = deps[:3]
        out[feature["id"]] = (", ".join(shown)
                              + (f" +{len(deps) - len(shown)}" if len(deps) > len(shown) else "")
                              + f" (clears W{max(1, -(-clears // 1)):.0f})")
    return out


def draw_gantt(workbook, option, epics, index, style, title=None, subtitle=None, waits=None):
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.hyperlink import Hyperlink

    title = (title or f"Gantt — {option['archetype'].title()}")[:31]
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False

    weeks = max(1, int(option["schedule"]["weeks"] + 0.999))
    labels = ["Epic / story", "Hours", "Owners", "Starts", "Ends", "Waits for"]
    first_week = len(labels) + 1

    sheet.cell(row=1, column=1, value=f"{option['archetype'].title()} — "
                                     f"{team_label(option['team_shape'])}").font = style["band_font"]
    phases = option.get("phases") or []
    lead = ""
    if len(phases) > 1:
        lead = (f"Phase {phases[0]['phase']} closes W"
                f"{max(1, -(-phases[0]['finishes_week'] // 1)):.0f} at "
                f"{phases[0]['standalone_hours']:.0f} h on its own · whole programme ")
    sheet.cell(row=2, column=1,
               value=(lead + f"{option['schedule']['weeks']:.1f} weeks · "
                      f"{option['estimate']['total_hours']['likely']:.0f} h likely · "
                      f"fit {option['score']['fit']:.1f}/10 · weeks are relative, "
                      f"week 1 is whenever this starts. A story's bar is the window it is "
                      f"OPEN, not the time it is worked — a story waiting between its build "
                      f"and its review is still inside its own bar. Expand a row for the "
                      f"component windows and who is on each.")).font = style["muted_font"]
    if subtitle:
        sheet.cell(row=3, column=1, value=subtitle).font = style["muted_font"]

    head = 4
    for i, label in enumerate(labels, start=1):
        cell = sheet.cell(row=head, column=i, value=label)
        cell.fill, cell.font = style["header_fill"], style["header_font"]
    for w in range(weeks):
        cell = sheet.cell(row=head, column=first_week + w, value=f"W{w + 1}")
        cell.fill, cell.font, cell.alignment = (style["header_fill"], style["header_font"],
                                                style["centre"])

    # The phase boundary, named on the row above the grid. A phase is a delivery commitment —
    # often a separate contract — so a reader has to be able to see where one ends without
    # counting bars, and the plan drew Phase 2 inside Phase 1 for exactly as long as nothing
    # marked it.
    if len(phases) > 1:
        # Column 1 of this row belongs to the subtitle when there is one — the baseline chart
        # uses it — so the label only takes the cell if it is free. The markers themselves sit
        # in week columns and never collide with it.
        if not subtitle:
            sheet.cell(row=head - 1, column=1,
                       value="Phase boundaries").font = style["muted_font"]
        for row_ in phases[:-1]:
            at = first_week + min(weeks - 1, max(0, int(row_["finishes_week"])))
            cell = sheet.cell(row=head - 1, column=at,
                              value=f"end P{row_['phase']}")
            cell.font, cell.alignment = style["muted_font"], style["centre"]

    def week_of(value):
        return f"W{max(1, -(-value // 1)):.0f}"

    row = head + 1
    for epic in story_rows(option, epics):
        # The epic is the GROUPING. Its band bounds its stories — earliest start to latest
        # finish — and its bar is a backdrop for them rather than work in its own right.
        cell = sheet.cell(row=row, column=1,
                          value=f"{epic['epic_id']}  {epic['name']}"
                          if epic["epic_id"] not in PSEUDO_EPIC else epic["name"])
        cell.fill, cell.font = style["band_fill"], style["band_font"]
        sheet.cell(row=row, column=2, value=round(epic["hours"], 1)).font = style["band_font"]
        sheet.cell(row=row, column=4,
                   value=f"W{int(epic['start_week']) + 1}").font = style["band_font"]
        sheet.cell(row=row, column=5,
                   value=week_of(epic["finish_week"])).font = style["band_font"]
        epic_gate = (waits or {}).get(epic["epic_id"])
        if epic_gate:
            sheet.cell(row=row, column=6, value=epic_gate).font = style["band_font"]
        for w in range(weeks):
            sheet.cell(row=row, column=first_week + w).fill = style["band_fill"]
        row += 1

        for story in epic["stories"]:
            sheet.row_dimensions[row].outlineLevel = 1
            label = (f"  {story['story_id']}  {story['label']}" if story["story_id"]
                     else f"  {COMPONENT_LABEL.get(story['slices'][0]['component'], story['label'])}")
            cell = sheet.cell(row=row, column=1, value=label)
            at = index.get(story["story_id"])
            if at:
                # Its OWN row on the Stories tab — the epic-packet version linked to whichever
                # story happened to be first in the packet. `display` is what Google Sheets
                # shows; without it the import turns this into a bare
                # `=HYPERLINK("#gid=...&range=A77")` and the reader gets the address.
                cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"Stories!A{at}",
                                           display=link_label(cell.value))
                cell.font = style["link_font"]
            else:
                cell.font = style["body_font"]
            sheet.cell(row=row, column=2, value=round(story["hours"], 1)).font = style["body_font"]
            sheet.cell(row=row, column=3,
                       value=", ".join(story["owners"])).font = style["muted_font"]
            sheet.cell(row=row, column=4,
                       value=f"W{int(story['start_week']) + 1}").font = style["body_font"]
            sheet.cell(row=row, column=5,
                       value=week_of(story["finish_week"])).font = style["body_font"]
            # A story answers with its OWN dependencies where it has any — more specific and
            # more useful than its epic's, because it names the actual work in the way. It falls
            # back to the epic's gate when it records none of its own.
            gate = (waits or {}).get(story["story_id"]) or epic_gate
            if gate:
                sheet.cell(row=row, column=6, value=gate).font = style["muted_font"]
            _paint(sheet, row, first_week,
                   span_weeks(story["start_week"], story["finish_week"], weeks),
                   style["role_fill"].get(story["role"]))
            row += 1

            for slice_ in story["slices"]:
                sheet.row_dimensions[row].outlineLevel = 2
                sheet.row_dimensions[row].hidden = True
                sheet.cell(row=row, column=1, value="      " + COMPONENT_LABEL.get(
                    slice_["component"], slice_["component"])).font = style["muted_font"]
                sheet.cell(row=row, column=2,
                           value=round(slice_["hours"], 1)).font = style["muted_font"]
                sheet.cell(row=row, column=3,
                           value=", ".join(sorted(slice_["owners"]))).font = style["muted_font"]
                sheet.cell(row=row, column=4,
                           value=f"W{int(slice_['start_week']) + 1}").font = style["muted_font"]
                sheet.cell(row=row, column=5,
                           value=week_of(slice_["finish_week"])).font = style["muted_font"]
                _paint(sheet, row, first_week,
                       span_weeks(slice_["start_week"], slice_["finish_week"], weeks),
                       style["role_fill"].get(story["role"]))
                row += 1

    # The two calendar-priced roles. Neither is on the team: the architect is `setup + a capped
    # weekly rate` so a second one cannot be priced, and overhead is ceremony charged per person
    # per week. Both are drawn from THIS option's own estimate and span the plan, because that
    # is exactly what they are — and leaving them off meant the chart accounted for about half
    # the hours the deal was being sold on, with the Architect band never drawn at all.
    for name, key in (("Architect", "architect"), ("Ceremony", "overhead")):
        component = (option["estimate"].get("project_components") or {}).get(key) or {}
        hours = component.get("hours") or 0.0
        if hours <= 0:
            continue
        cell = sheet.cell(row=row, column=1, value=name)
        cell.fill, cell.font = style["band_fill"], style["band_font"]
        sheet.cell(row=row, column=2, value=round(hours, 1)).font = style["band_font"]
        sheet.cell(row=row, column=4, value="W1").font = style["band_font"]
        sheet.cell(row=row, column=5, value=f"W{weeks}").font = style["band_font"]
        _paint(sheet, row, first_week, range(weeks), style["role_fill"].get(key)
               or style["band_fill"])
        row += 1
        sheet.row_dimensions[row].outlineLevel = 1
        sheet.cell(row=row, column=1, value=(
            f"  {'setup plus a capped weekly rate, across the plan' if key == 'architect' else 'ceremony, per person per week'}"
        )).font = style["muted_font"]
        row += 1

    # Reconciliation. The chart and the estimate are two views of one number and a reader is
    # entitled to check that they agree — three different dev figures once sat in one workbook
    # with nothing saying which was which.
    scheduled = sum(p["delivered_hours"] for p in option["schedule"]["team"])
    calendar_priced = sum((option["estimate"]["project_components"].get(k) or {}).get("hours", 0.0)
                          for k in ("architect", "overhead"))
    priced = sum(r["hours"] for r in option["estimate"]["by_role"].values())
    row += 1
    sheet.cell(row=row, column=1, value="Drawn above").font = style["band_font"]
    sheet.cell(row=row, column=2, value=round(scheduled + calendar_priced, 1)).font = style["band_font"]
    sheet.cell(row=row, column=first_week, value=(
        f"Scheduled {scheduled:.0f} h + architect and ceremony {calendar_priced:.0f} h "
        f"= {scheduled + calendar_priced:.0f} h, against {priced:.0f} h on the estimate "
        f"({(scheduled + calendar_priced) / priced:.0%}). Hours here are expected values (PERT "
        f"means); the Stories tab quotes the mode of each interval, which is a different and "
        f"slightly smaller figure.")).font = style["muted_font"]
    row += 1

    sheet.freeze_panes = sheet.cell(row=head + 1, column=first_week)
    sheet.column_dimensions["A"].width = 46
    for i in range(2, first_week):
        sheet.column_dimensions[get_column_letter(i)].width = 9
    for w in range(weeks):
        sheet.column_dimensions[get_column_letter(first_week + w)].width = 4.2
    sheet.sheet_properties.outlinePr.summaryBelow = False
    return sheet


def _paint(sheet, row, first_week, columns, fill):
    """Fill the given week columns on a row."""
    if fill is None:
        return
    for w in sorted(columns):
        sheet.cell(row=row, column=first_week + w).fill = fill


def draw_options(workbook, plan, style):
    """Every option side by side: the schedule, the score, and the role hours it costs."""
    title = "Options"
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False

    sheet.cell(row=1, column=1, value=f"Delivery options — {plan.get('project') or ''}"
               ).font = style["band_font"]
    sheet.cell(row=2, column=1, value=plan["estimate"]["why_it_differs"]).font = style["muted_font"]

    columns = (["Option", "Archetype", "Team", "Weeks", "Likely h", "Low", "High", "Fit",
                "Deliverable"] + [ROLE_LABEL.get(r, r) for r in ROLE_ORDER]
               + ["Drift h", "Utilisation"])
    head = 4
    for i, name in enumerate(columns, start=1):
        cell = sheet.cell(row=head, column=i, value=name)
        cell.fill, cell.font = style["header_fill"], style["header_font"]

    for n, o in enumerate(plan["options"], start=1):
        row = head + n
        t = o["estimate"]["total_hours"]
        util = [p["utilisation"] for p in o["schedule"]["team"]]
        values = ([("★ " if o["id"] == plan.get("recommended") else "") + o["id"],
                   o["archetype"], team_label(o["team_shape"]),
                   round(o["schedule"]["weeks"], 1), round(t["likely"]), round(t["low"]),
                   round(t["high"]), o["score"]["fit"],
                   "yes" if o["feasibility"]["feasible"] else "NO"]
                  + [round(o["estimate"]["by_role"].get(r, {}).get("hours", 0.0))
                     for r in ROLE_ORDER]
                  + [o["schedule"]["drift_hours"],
                     round(sum(util) / len(util), 2) if util else 0.0])
        for i, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=i, value=value)
            cell.font = style["band_font"] if o["id"] == plan.get("recommended") else style["body_font"]
            cell.border = style["rule"]

    sheet.freeze_panes = f"A{head + 1}"
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["C"].width = 34
    from openpyxl.utils import get_column_letter
    for i in range(4, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(i)].width = 11
    return sheet


def extend_workbook(plan, path, per="archetype"):
    """Append the Options tab and the Gantts to the estimate's own workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return False, "openpyxl is not importable"
    path = Path(path)
    if not path.exists():
        return False, f"{path} does not exist — run est-estimate's render first"

    style = brand_module().xlsx()
    workbook = load_workbook(path)
    index = story_index(workbook)
    draw_options(workbook, plan, style)

    # One Gantt per SCHEDULE OFFERED, which is the best team shape for each archetype — not one
    # per row of the sweep. Nine near-identical charts is not a more readable document than
    # three, and readability is the requirement; the Options tab carries every row.
    #
    # The BASELINE comes first, always. Picking the best-fit shape of each archetype always
    # picks the largest team the sweep allowed — all three charts in one real workbook were
    # `2x BA, 4x Dev, 1x DevOps, 1x QA, 1x UX` — so the one-person-per-role reading, which is
    # what a reader calibrates every other option against, was the single schedule the workbook
    # did not contain.
    drawn = []
    if per == "option":
        chosen = list(plan["options"])
    else:
        chosen = []
        for name in ("sequential", "foundation", "pipelined"):
            same = [o for o in plan["options"] if o["archetype"] == name]
            if same:
                chosen.append(max(same, key=lambda o: o["score"]["fit"]))

    # Story-level dependencies live on the priced features, which every option carries a copy
    # of. Taken from whichever option is to hand: the scope is one fact and re-pricing a
    # calendar never moves a story's `depends_on`.
    sample = plan.get("baseline") or (plan.get("options") or [{}])[0]
    features = (sample.get("estimate") or {}).get("features") or []

    baseline = plan.get("baseline")
    if baseline is not None:
        best = next((o for o in plan["options"] if o["id"] == plan.get("recommended")), None)
        against = (f" Compare against {best['archetype'].title()}, {team_label(best['team_shape'])}: "
                   f"{best['schedule']['weeks']:.1f} weeks, "
                   f"{best['estimate']['total_hours']['likely']:.0f} h."
                   if best is not None and best["id"] != baseline["id"] else "")
        draw_gantt(workbook, baseline, plan.get("epics") or [], index, style,
                   title="Gantt — Baseline (1 each)",
                   subtitle=baseline.get("baseline_note", "") + against,
                   waits=gates(plan, baseline, features))
        drawn.append(baseline["id"] + " (baseline)")
        chosen = [o for o in chosen if o["id"] != baseline["id"]]

    for option in chosen:
        draw_gantt(workbook, option, plan.get("epics") or [], index, style,
                   waits=gates(plan, option, features))
        drawn.append(option["id"])
    workbook.save(path)
    return True, drawn


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("plan", help="path to plan.json")
    ap.add_argument("-o", "--output", help="write plan.md here (default: beside plan.json)")
    ap.add_argument("--workbook", help="estimate.xlsx to extend with the Gantt tabs")
    ap.add_argument("--gantt-per", choices=("archetype", "option"), default="archetype",
                    help="a Gantt for the best shape of each archetype (default) or for every option")
    ap.add_argument("--no-markdown", action="store_true")
    args = ap.parse_args()

    plan = json.loads(Path(args.plan).read_text(encoding="utf-8"))
    written, notes = [], []
    if not args.no_markdown:
        target = Path(args.output or Path(args.plan).with_name("plan.md"))
        target.write_text(markdown(plan), encoding="utf-8")
        written.append(str(target))
    if args.workbook:
        ok, result = extend_workbook(plan, args.workbook, args.gantt_per)
        if ok:
            written.append(str(args.workbook))
            notes.append(f"gantt tabs: {', '.join(result)}")
        else:
            notes.append(f"workbook skipped: {result}")
    print(json.dumps({"ok": True, "written": written, "notes": notes}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
