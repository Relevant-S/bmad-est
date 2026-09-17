#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Tests for render-plan.py — the markdown, and the Gantt tabs appended to the estimate's own
workbook.

The workbook round trip is the risk. est-plan opens a file another skill wrote and saves it
again, so the question that matters is not whether the Gantt is correct but whether the
Stories and Tasks tabs, and the hyperlinks between them, are still there afterwards. A plan
that quietly strips the estimate out of the estimate workbook would be discovered by a client.
"""

import collections
import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import fixtures as F  # noqa: E402

render = F.load("render_plan", F.PLAN / "render-plan.py")
render_estimate = F.load("render_estimate", F.EST / "scripts" / "render-estimate.py")

try:
    import openpyxl  # noqa: F401
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False


def a_plan(features=None):
    estimate = F.priced(features or F.wide(18, epics=3, size="L"))
    return estimate, F.plan_mod.build_plan(estimate, F.model())


class TheMarkdown(unittest.TestCase):
    def setUp(self):
        self.estimate, self.plan = a_plan()
        self.text = render.markdown(self.plan)

    def test_every_option_is_presented_with_its_own_hours_not_just_its_calendar(self):
        """A schedule without its cost is half an answer. The module's rule is that the reader
        compares plan-and-cost pairs, never calendars."""
        for option in self.plan["options"]:
            self.assertIn(f"{option['estimate']['total_hours']['likely']:.0f}", self.text)
        self.assertIn("| Role | Low | Likely | High |", self.text)

    def test_the_sub_scores_are_printed_not_just_the_total(self):
        self.assertIn("| Sub-score | /10 | What it rests on |", self.text)
        for name in ("Duration", "Cost", "Utilisation", "Coordination", "Drift", "Confidence"):
            self.assertIn(f"| {name} |", self.text)

    def test_every_headcount_decision_appears_with_its_reasoning(self):
        for decision in self.plan["staffing"]["decisions"]:
            needle = (decision["justification"] or decision["refusals"])[0].strip()
            self.assertIn(needle[:60], self.text)

    def test_each_option_names_its_risks_and_what_to_do_about_them(self):
        self.assertIn("**Risks**", self.text)
        for option in self.plan["options"]:
            for row in option["risks"]:
                self.assertIn(row["mitigation"][:50], self.text)

    def test_it_says_weeks_are_relative_and_carries_no_dates(self):
        import re
        self.assertIn("Weeks are relative", self.text)
        self.assertIsNone(re.search(r"\b\d{4}-\d{2}-\d{2}\b", self.text))

    def test_the_markdown_carries_the_build_order_and_its_evidence(self):
        """Every edge prints the reason it exists — the inventory's own sentence where there is
        one, and the story pair that caused it where there is not."""
        estimate = F.priced(F.layered(per=4, layers=4))
        plan = F.plan_mod.build_plan(estimate, F.model())
        text = render.markdown(plan)
        self.assertIn("## Build order", text)
        edges = [w for e in plan["build_order"]["epics"] for w in e["waits_for"]]
        self.assertTrue(edges, "the fixture recorded no ordering to print")
        for wait in edges:
            self.assertIn(wait["why"][:40], text)

    def test_a_plan_whose_calendar_contradicts_its_graph_says_so_loudly(self):
        plan = json.loads(json.dumps(self.plan))
        plan["ordering_check"] = {"ok": False, "findings": [
            {"kind": "epic_build_order",
             "detail": "E4 is built before E3, invented for the test"}]}
        text = render.markdown(plan)
        self.assertIn("contradicts its own dependency graph", text)
        self.assertIn("invented for the test", text)
        self.assertIn("not shippable", text)

    def test_an_infeasible_option_is_marked_in_the_table_not_quietly_listed(self):
        plan = json.loads(json.dumps(self.plan))
        plan["options"][0]["feasibility"] = {"feasible": False, "checks": [],
                                             "failures": ["invented for the test"],
                                             "critical_path_weeks": 0.0}
        plan["recommended"] = None
        text = render.markdown(plan)
        self.assertIn("not deliverable", text)
        self.assertIn("No option is deliverable", text)


@unittest.skipUnless(HAS_XLSX, "openpyxl is not installed")
class TheWorkbookRoundTrip(unittest.TestCase):
    def setUp(self):
        self.estimate, self.plan = a_plan()
        self.dir = tempfile.TemporaryDirectory()
        self.book = Path(self.dir.name) / "estimate.xlsx"
        self.assertTrue(render_estimate.write_xlsx(self.estimate, self.book))
        self.addCleanup(self.dir.cleanup)

    def extend(self, per="archetype"):
        ok, result = render.extend_workbook(self.plan, self.book, per)
        self.assertTrue(ok, result)
        from openpyxl import load_workbook
        return load_workbook(self.book)

    def test_the_estimates_own_tabs_survive_being_extended(self):
        """The one thing that must not break. This opens a file est-estimate wrote."""
        from openpyxl import load_workbook
        before = load_workbook(self.book)
        rows, cols = before["Stories"].max_row, before["Stories"].max_column
        after = self.extend()
        self.assertEqual(after["Stories"].max_row, rows)
        self.assertEqual(after["Stories"].max_column, cols)
        self.assertIn("Tasks", after.sheetnames)
        self.assertEqual(after["Stories"].freeze_panes, before["Stories"].freeze_panes)

    def test_it_appends_rather_than_replacing(self):
        after = self.extend()
        self.assertEqual(after.sheetnames[:2], ["Stories", "Tasks"])
        self.assertIn("Options", after.sheetnames)
        self.assertTrue([n for n in after.sheetnames if n.startswith("Gantt")])

    def test_one_gantt_per_archetype_by_default_and_one_per_option_on_request(self):
        """Nine near-identical charts is not a more readable document than three, and
        readability is a requirement rather than a preference.

        Plus the baseline, which is always drawn and is never one of the per-archetype charts
        unless it happens to be the best-fit shape for its own archetype.
        """
        names = [n for n in self.extend().sheetnames if n.startswith("Gantt")]
        archetypes = {o["archetype"] for o in self.plan["options"]}
        baseline = self.plan.get("baseline")
        extra = 1 if baseline is not None and baseline["id"] not in {
            max((o for o in self.plan["options"] if o["archetype"] == a),
                key=lambda o: o["score"]["fit"])["id"] for a in archetypes} else 0
        self.assertEqual(len(names), len(archetypes) + extra)
        every = self.extend("option")
        self.assertGreaterEqual(len([n for n in every.sheetnames if n.startswith("Gantt")]), 1)

    def test_rendering_twice_does_not_duplicate_the_tabs(self):
        self.extend()
        after = self.extend()
        self.assertEqual(len(after.sheetnames), len(set(after.sheetnames)))

    def test_every_story_row_on_a_gantt_links_back_to_the_stories_tab(self):
        after = self.extend()
        sheet = after[[n for n in after.sheetnames if n.startswith("Gantt")][0]]
        index = render.story_index(after)
        self.assertTrue(index)
        links = [c for row in sheet.iter_rows() for c in row if c.hyperlink]
        self.assertTrue(links, "no Gantt row linked to its work")
        for cell in links:
            self.assertTrue(cell.hyperlink.location.startswith("Stories!A"))
            target = int(cell.hyperlink.location.split("A")[1])
            self.assertIn(target, set(index.values()))

    def test_the_grid_has_one_column_per_week_and_freezes_the_labels(self):
        after = self.extend()
        name = next(n for n in after.sheetnames
                    if any(f"Gantt — {o['archetype'].title()}"[:31] == n
                           for o in self.plan["options"]))
        sheet, option = after[name], None
        for o in self.plan["options"]:
            if f"Gantt — {o['archetype'].title()}"[:31] == name:
                option = o
                break
        weeks = max(1, int(option["schedule"]["weeks"] + 0.999))
        # Six label columns. Rows are STORIES grouped under epics, so the first column names
        # the work rather than the worker and `Owners` carries the people. `Joins` is gone with
        # the person rows; the arrival table in plan.md is where that lives now.
        self.assertEqual([sheet.cell(row=4, column=c).value for c in range(1, 7)],
                         ["Epic / story", "Hours", "Owners", "Starts", "Ends", "Waits for"])
        self.assertEqual(sheet.cell(row=4, column=7).value, "W1")
        self.assertEqual(sheet.cell(row=4, column=6 + weeks).value, f"W{weeks}")
        self.assertEqual(sheet.freeze_panes, "G5")

    def test_the_rows_are_epics_then_stories_then_component_slices(self):
        """Epics are the GROUPING; the rows are the stories. The first version drew one row per
        person per epic and rolled the stories into a count, so the finest thing a reader could
        see was an epic — and a schedule is a statement about work items."""
        sheet = self.extend()[[n for n in self.extend().sheetnames if n.startswith("Gantt")][0]]
        levels = {sheet.row_dimensions[r].outlineLevel for r in range(5, sheet.max_row + 1)}
        self.assertEqual(levels, {0, 1, 2})
        # Level 2 is detail on demand, so it starts collapsed.
        hidden = [sheet.row_dimensions[r].hidden for r in range(5, sheet.max_row + 1)
                  if sheet.row_dimensions[r].outlineLevel == 2]
        self.assertTrue(hidden and all(hidden))

    def test_every_scheduled_story_is_drawn_exactly_once(self):
        """A story appears once however many people touch it. Counted against the schedule's
        own bookings, so a story silently dropped from the chart is caught."""
        after = self.extend()
        option = max(self.plan["options"], key=lambda o: o["score"]["fit"])
        sheet = after[f"Gantt — {option['archetype'].title()}"[:31]]
        want = {i["story_id"] for p in option["schedule"]["team"] for i in p["items"]
                if i.get("story_id")}
        drawn = collections.Counter()
        for r in range(5, sheet.max_row + 1):
            if sheet.row_dimensions[r].outlineLevel != 1:
                continue
            label = str(sheet.cell(row=r, column=1).value or "").strip()
            drawn[label.split("  ")[0]] += 1
        self.assertTrue(want)
        for story in want:
            self.assertEqual(drawn.get(story), 1,
                             f"{story} is drawn {drawn.get(story)} times, not once")

    def test_a_story_links_to_its_own_row_not_a_neighbours(self):
        """The epic-packet version linked to whichever story happened to be first in the
        packet, so four of five rows pointed at somebody else's work."""
        after = self.extend()
        index = render.story_index(after)
        option = max(self.plan["options"], key=lambda o: o["score"]["fit"])
        sheet = after[f"Gantt — {option['archetype'].title()}"[:31]]
        checked = 0
        for r in range(5, sheet.max_row + 1):
            cell = sheet.cell(row=r, column=1)
            if sheet.row_dimensions[r].outlineLevel != 1 or not cell.hyperlink:
                continue
            story = str(cell.value or "").strip().split("  ")[0]
            self.assertEqual(cell.hyperlink.location, f"Stories!A{index[story]}",
                             f"{story} links somewhere other than its own row")
            checked += 1
        self.assertGreater(checked, 0, "no story row carried a link")

    def test_a_story_expands_to_its_component_windows_and_their_owners(self):
        after = self.extend()
        option = max(self.plan["options"], key=lambda o: o["score"]["fit"])
        sheet = after[f"Gantt — {option['archetype'].title()}"[:31]]
        booked = collections.defaultdict(lambda: collections.defaultdict(set))
        for person in option["schedule"]["team"]:
            for item in person["items"]:
                if item.get("story_id"):
                    booked[item["story_id"]][item["component"]].add(person["name"])
        story, slices = None, collections.defaultdict(set)
        found = 0
        for r in range(5, sheet.max_row + 1):
            depth = sheet.row_dimensions[r].outlineLevel
            label = str(sheet.cell(row=r, column=1).value or "").strip()
            if depth == 1:
                if story in booked and slices:
                    found += 1
                    for name, owners in slices.items():
                        self.assertEqual(owners, booked[story][name], f"{story} {name}")
                story, slices = label.split("  ")[0], collections.defaultdict(set)
            elif depth == 2 and story:
                key = {v: k for k, v in render.COMPONENT_LABEL.items()}.get(label)
                if key:
                    slices[key] = set(str(sheet.cell(row=r, column=3).value or "").split(", "))
        self.assertGreater(found, 3, "no story expanded to its components")

    def test_the_header_says_the_bar_is_elapsed_time_not_occupancy(self):
        """A story sitting between its build and its review is inside its own bar. Saying so is
        the whole reason it is safe to draw one solid bar per story."""
        after = self.extend()
        sheet = after[[n for n in after.sheetnames if n.startswith("Gantt")][0]]
        self.assertIn("not the time it is worked", sheet.cell(row=2, column=1).value)

    def test_an_epic_band_bounds_the_stories_under_it(self):
        after = self.extend()
        option = max(self.plan["options"], key=lambda o: o["score"]["fit"])
        sheet = after[f"Gantt — {option['archetype'].title()}"[:31]]
        band, starts, ends, checked = None, [], [], 0
        def close():
            nonlocal checked
            # Architect and Ceremony are calendar-priced bands with no stories under them —
            # their sub-row is a note, not work — so there is nothing to bound.
            if band and starts and all(starts) and all(ends):
                self.assertEqual(band[0], min(starts), f"{band[2]} starts after its first story")
                self.assertEqual(band[1], max(ends), f"{band[2]} ends before its last story")
                checked += 1
        for r in range(5, sheet.max_row + 1):
            depth = sheet.row_dimensions[r].outlineLevel
            cells = [sheet.cell(row=r, column=c).value for c in (1, 4, 5)]
            if depth == 0 and cells[1]:
                close()
                band, starts, ends = (cells[1], cells[2], cells[0]), [], []
            elif depth == 1 and band:
                starts.append(cells[1])
                ends.append(cells[2])
        close()
        self.assertGreater(checked, 2, "no epic band was checked against its stories")

    def test_bars_are_painted_in_the_brand_colour_for_the_role(self):
        after = self.extend()
        sheet = after[[n for n in after.sheetnames if n.startswith("Gantt")][0]]
        brand = F.load("brand", F.EST / "scripts" / "brand.py")
        wanted = {brand._hex(c) for r, c in brand.load()["roles"].items() if not r.startswith("_")}
        painted = {sheet.cell(row=r, column=c).fill.fgColor.rgb
                   for r in range(5, sheet.max_row + 1) for c in range(7, sheet.max_column + 1)}
        painted = {p[2:] if isinstance(p, str) and len(p) == 8 else p for p in painted}
        self.assertTrue(wanted & painted, "no bar used a role colour from brand.json")

    def test_the_options_tab_carries_every_option_and_marks_the_recommendation(self):
        sheet = self.extend()["Options"]
        ids = [sheet.cell(row=r, column=1).value for r in range(5, sheet.max_row + 1)]
        self.assertEqual(len([i for i in ids if i]), len(self.plan["options"]))
        self.assertTrue(any(str(i).startswith("★") for i in ids))

    def test_hours_reconcile_from_slice_to_story_to_epic(self):
        """Three levels of the same number. A chart whose own rows do not add up is not
        evidence of anything."""
        after = self.extend()
        option = max(self.plan["options"], key=lambda o: o["score"]["fit"])
        sheet = after[f"Gantt — {option['archetype'].title()}"[:31]]
        epic = story = None
        epic_total = story_total = slice_total = 0.0
        checked = 0
        rows = list(range(5, sheet.max_row + 1)) + [None]
        for r in rows:
            depth = sheet.row_dimensions[r].outlineLevel if r else 0
            hours = (sheet.cell(row=r, column=2).value or 0.0) if r else 0.0
            label = str(sheet.cell(row=r, column=1).value or "") if r else ""
            if depth == 2:
                slice_total += hours
                continue
            if story is not None:
                self.assertAlmostEqual(slice_total, story, delta=0.15, msg="slices vs story")
                checked += 1
            if depth == 1:
                story, slice_total = hours, 0.0
                story_total += hours
                continue
            story = None
            if epic is not None and label != "Drawn above":
                self.assertAlmostEqual(story_total, epic, delta=0.2, msg="stories vs epic")
            epic, story_total, slice_total = hours, 0.0, 0.0
            if label.startswith(("Architect", "Ceremony", "Drawn above")):
                epic = None
        self.assertGreater(checked, 5)

    def test_the_calendar_priced_roles_are_drawn_at_all(self):
        """The architect is on no team — it is priced as setup plus a capped weekly rate, so a
        second one cannot be priced — and ceremony is charged per person per week. Neither was
        ever drawn, which is most of why the chart accounted for about half the hours."""
        after = self.extend()
        sheet = after[[n for n in after.sheetnames if n.startswith("Gantt")][0]]
        labels = [sheet.cell(row=r, column=1).value for r in range(5, sheet.max_row + 1)]
        self.assertIn("Architect", labels)
        self.assertIn("Ceremony", labels)

    def test_the_chart_reconciles_with_the_estimate_it_belongs_to(self):
        """Three different dev figures once sat in one workbook with nothing saying which was
        which. The chart and the estimate are two views of one number."""
        after = self.extend()
        sheet = after[[n for n in after.sheetnames if n.startswith("Gantt")][0]]
        note = None
        for r in range(5, sheet.max_row + 1):
            if sheet.cell(row=r, column=1).value == "Drawn above":
                note = sheet.cell(row=r, column=7).value
                drawn = sheet.cell(row=r, column=2).value
        self.assertIsNotNone(note, "no reconciliation row on the chart")
        self.assertIn("against", note)
        self.assertIn("PERT", note)
        # The FIRST Gantt is the baseline when there is one, so reconcile against that option
        # rather than against the best-scoring one.
        title = [n for n in after.sheetnames if n.startswith("Gantt")][0]
        option = (self.plan.get("baseline")
                  if title.startswith("Gantt — Baseline")
                  else max(self.plan["options"], key=lambda o: o["score"]["fit"]))
        priced = sum(r["hours"] for r in option["estimate"]["by_role"].values())
        self.assertGreater(drawn / priced, 0.9,
                           "the chart still leaves a tenth of the priced hours undrawn")

    def test_the_baseline_chart_is_drawn_first_and_holds_one_person_per_role(self):
        """Picking the best-fit shape of each archetype always picks the LARGEST team the sweep
        allowed — all three Gantt tabs in one real workbook were `2x BA, 4x Dev, 1x DevOps,
        1x QA, 1x UX`. The one-per-role reading, which is what a reader calibrates the others
        against, was the single schedule the workbook did not contain."""
        after = self.extend()
        gantts = [n for n in after.sheetnames if n.startswith("Gantt")]
        self.assertEqual(gantts[0], "Gantt — Baseline (1 each)")
        baseline = self.plan["baseline"]
        self.assertTrue(all(n == 1 for n in baseline["team_shape"].values()),
                        baseline["team_shape"])
        self.assertEqual(len(baseline["schedule"]["team"]), len(baseline["team_shape"]))
        # It is an OPTION, so it carries its own estimate like every other one.
        self.assertGreater(baseline["estimate"]["total_hours"]["likely"], 0)
        self.assertIn(baseline["baseline_note"][:40],
                      after[gantts[0]].cell(row=3, column=1).value)

    def test_the_baseline_ignores_a_roster_that_is_already_larger(self):
        """It is the reading every other option is compared against, so it cannot move with
        whatever team happens to exist — that is the thing it is there to price against."""
        plan = F.plan_mod.build_plan(self.estimate, F.model(), roster={"dev": 2})
        baseline = plan["baseline"]
        self.assertEqual(baseline["team_shape"]["dev"], 1)
        self.assertIn("supplied", baseline["baseline_note"])
        # A roster is a floor: the sweep may add to it and must never propose fewer people than
        # the user already has. The baseline is a reference chart, so it stays off the options
        # list rather than offering them a team smaller than the one they told us they have.
        self.assertTrue(all(o["team_shape"].get("dev", 0) >= 2 for o in plan["options"]))

    def test_every_hyperlink_carries_the_display_text_google_sheets_needs(self):
        """Excel shows the cell's own value, which is why this went unnoticed. Google Sheets
        rewrites the link into `=HYPERLINK("#gid=...&range=A77")` and uses `display` as the
        label — with none set it shows the address, and a column of epic names arrived reading
        `#gid=1123955261&range=A2`. The `gid` cannot be written from here, so `display` is the
        only lever."""
        after = self.extend()
        links = [(n, c) for n in after.sheetnames for row in after[n].iter_rows()
                 for c in row if c.hyperlink]
        self.assertTrue(links, "no links in the workbook at all")
        for name, cell in links:
            self.assertEqual(getattr(cell.hyperlink, "display", None), str(cell.value),
                             f"{name}!{cell.coordinate} would show its address in Sheets")

    def test_the_chart_says_what_each_epic_waits_for(self):
        """The plan has to answer where a piece of work belongs and on what grounds, off the
        chart rather than by reverse-engineering the bars.

        On an ordered backlog: the default fixture has no dependencies at all, so it has no
        build order to draw and would let this pass without drawing anything.
        """
        estimate = F.priced(F.layered(per=4, layers=4))
        plan = F.plan_mod.build_plan(estimate, F.model())
        book = Path(self.dir.name) / "ordered.xlsx"
        self.assertTrue(render_estimate.write_xlsx(estimate, book))
        ok, _ = render.extend_workbook(plan, book, "archetype")
        self.assertTrue(ok)
        from openpyxl import load_workbook
        after = load_workbook(book)
        sheet = after[[n for n in after.sheetnames if n.startswith("Gantt")][-1]]
        self.assertEqual(sheet.cell(row=4, column=6).value, "Waits for")
        bands = [sheet.cell(row=r, column=6).value
                 for r in range(5, sheet.max_row + 1)
                 if sheet.row_dimensions[r].outlineLevel == 0 and sheet.cell(row=r, column=6).value]
        self.assertTrue(bands, "no epic band names what it stands on")
        for gate in bands:
            self.assertRegex(gate, r"^E\S+")
        # And a story answers with its OWN dependencies where it has them, which is both more
        # specific and more useful than its epic's — it names the work actually in the way.
        stories = [sheet.cell(row=r, column=6).value
                   for r in range(5, sheet.max_row + 1)
                   if sheet.row_dimensions[r].outlineLevel == 1
                   and sheet.cell(row=r, column=6).value]
        self.assertTrue(any(g.startswith("F") for g in stories),
                        "no story row named its own dependency")

    def test_the_chart_marks_the_phase_boundary_and_leads_with_phase_one(self):
        """A phase is a delivery commitment and often a separate contract, so a reader must be
        able to see where one ends without counting bars — and the headline must not hand a
        client a span that silently includes work under another agreement."""
        features, epics = F.two_phases()
        estimate = F.priced(features, epics=epics)
        plan = F.plan_mod.build_plan(estimate, F.model())
        book = Path(self.dir.name) / "phased.xlsx"
        self.assertTrue(render_estimate.write_xlsx(estimate, book))
        ok, _ = render.extend_workbook(plan, book, "archetype")
        self.assertTrue(ok)
        from openpyxl import load_workbook
        sheet = load_workbook(book)[[n for n in load_workbook(book).sheetnames
                                     if n.startswith("Gantt")][-1]]
        self.assertIn("Phase 1 closes W", sheet.cell(row=2, column=1).value)
        self.assertIn("whole programme", sheet.cell(row=2, column=1).value)
        markers = [sheet.cell(row=3, column=c).value for c in range(1, sheet.max_column + 1)]
        self.assertIn("end P1", markers)

    def test_the_markdown_leads_with_phase_one_and_prints_both_figures(self):
        features, epics = F.two_phases()
        plan = F.plan_mod.build_plan(F.priced(features, epics=epics), F.model())
        text = render.markdown(plan)
        self.assertIn("**Phase 1:", text)
        self.assertIn("on its own", text)
        self.assertIn("Whole programme", text)
        # Both figures, and the sentence saying which to quote to a client — the apportioned
        # number is not what anyone saves by dropping a phase.
        self.assertIn("| Phase | Runs | On its own | Share of this programme |", text)
        self.assertIn("paid once however many", text)
        self.assertIn("| P1 weeks | P1 h | Whole weeks | Whole h |", text)

    def test_a_single_phase_backlog_reads_exactly_as_before(self):
        """The reporting change must be inert where there is only one phase."""
        text = render.markdown(self.plan)
        self.assertNotIn("**Phase 1:", text)
        self.assertIn("| # | Option | Weeks | Likely h | Range | Fit |", text)

    def test_a_missing_workbook_is_reported_rather_than_created(self):
        """est-plan extends the estimate's workbook. Creating one would hand a client a second
        file with a Gantt and no scope in it."""
        ok, why = render.extend_workbook(self.plan, Path(self.dir.name) / "absent.xlsx")
        self.assertFalse(ok)
        self.assertIn("does not exist", why)


if __name__ == "__main__":
    unittest.main()
